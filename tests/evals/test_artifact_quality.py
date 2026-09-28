"""Exercise real fixture outputs and reject misleading benchmark passes."""

from __future__ import annotations

import copy
import json
from io import BytesIO
from pathlib import Path
from typing import Any, cast
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document

from evals import artifact_quality
from evals import artifact_quality_documents as documents
from maivn_tools import PDFToolSet
from maivn_tools.connectors.workbooks import toolset as workbook_toolset


@pytest.fixture(scope='module')
def benchmark(tmp_path_factory: pytest.TempPathFactory) -> dict[str, Any]:
    return artifact_quality.run(tmp_path_factory.mktemp('artifact-quality'))


def test_actual_format_evaluators_pass_without_claiming_visual_or_model_review(
    benchmark: dict[str, Any],
) -> None:
    assert benchmark['automated']['status'] == 'pass'
    assert not artifact_quality.validate_report(benchmark)
    assert set(benchmark['formats']) == {'docx', 'pdf', 'pptx', 'xlsx'}
    assert benchmark['live_model']['status'] == 'not_run'
    for result in benchmark['formats'].values():
        assert result['automated']['status'] == 'pass'
        assert result['native_render']['status'] == 'pending'
        assert result['visual_review']['status'] == 'pending'
        for item in result['artifacts'].values():
            assert Path(item['path']).is_file()
    assert benchmark['formats']['xlsx']['calculation']['status'] == 'pending'


@pytest.mark.parametrize('checks', [{}, {'missing': False}, {'truthy': 1}])
def test_empty_false_and_non_boolean_gates_fail_closed(
    benchmark: dict[str, Any],
    checks: dict[str, object],
) -> None:
    altered = copy.deepcopy(benchmark)
    altered['formats']['pptx']['automated']['checks'] = checks
    assert artifact_quality.validate_report(altered)


def test_tampered_file_fails_hash_binding(benchmark: dict[str, Any], tmp_path: Path) -> None:
    altered = copy.deepcopy(benchmark)
    revised = altered['formats']['pptx']['artifacts']['revised']
    damaged = tmp_path / 'tampered.pptx'
    damaged.write_bytes(Path(revised['path']).read_bytes() + b'tampered')
    revised['path'] = str(damaged)
    assert any('sha256' in error for error in artifact_quality.validate_report(altered))


def test_missing_format_cannot_pass(benchmark: dict[str, Any]) -> None:
    altered = copy.deepcopy(benchmark)
    del altered['formats']['pdf']
    assert artifact_quality.validate_report(altered)


def test_formula_deletion_causes_a_failed_integrated_report(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    original = workbook_toolset.render_xlsx_bytes

    def strip_formulas(manifest: dict[str, object]) -> bytes:
        target = BytesIO()
        with (
            ZipFile(BytesIO(original(manifest))) as source,
            ZipFile(
                target,
                'w',
                ZIP_DEFLATED,
            ) as output,
        ):
            for name in source.namelist():
                content = source.read(name)
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    # The XML is generated immediately above by our fixture renderer.
                    sheet = ET.fromstring(content)  # noqa: S314
                    for cell in sheet.iter():
                        for child in list(cell):
                            if child.tag.rsplit('}', 1)[-1] == 'f':
                                cell.remove(child)
                    content = ET.tostring(sheet)
                output.writestr(name, content)
        return target.getvalue()

    monkeypatch.setattr(workbook_toolset, 'render_xlsx_bytes', strip_formulas)
    report = artifact_quality.run(tmp_path)
    assert report['automated']['status'] == 'fail'
    assert report['formats']['xlsx']['automated']['status'] == 'fail'
    assert report['formats']['pptx']['automated']['status'] == 'pass'


def test_cli_requires_output_and_exits_nonzero_for_failed_report(
    benchmark: dict[str, Any],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    with pytest.raises(SystemExit) as missing:
        artifact_quality.main([])
    assert missing.value.code == 2
    altered = copy.deepcopy(benchmark)
    altered['formats']['docx']['automated']['checks'] = {}

    def failed_report(_: Path) -> dict[str, Any]:
        return altered

    monkeypatch.setattr(artifact_quality, 'run', failed_report)
    assert artifact_quality.main(['--output-dir', str(tmp_path)]) == 1
    saved = cast('dict[str, Any]', json.loads((tmp_path / 'report.json').read_text()))
    assert saved['automated']['status'] == 'fail'


def test_document_fixture_has_the_exact_package_sequence(benchmark: dict[str, Any]) -> None:
    path = benchmark['formats']['docx']['artifacts']['initial']['path']
    rows = Document(path).tables[0].rows
    assert [row.cells[0].text for row in rows] == [
        'Package',
        *[f'WP-{index:02d}' for index in range(1, 29)],
    ]
    assert [row.cells[2].text for row in rows[1:]] == [
        *(['Delivery', 'Service', 'Operations'] * 9),
        'Delivery',
    ]


@pytest.mark.parametrize(
    ('column', 'replacement'), [(0, 'WP-01'), (1, 'Evidence removed'), (2, 'Unknown')]
)
@pytest.mark.parametrize('extension', ['docx', 'pdf'])
def test_document_record_mutations_fail_completeness(
    benchmark: dict[str, Any],
    tmp_path: Path,
    extension: str,
    column: int,
    replacement: str,
) -> None:
    document = Document(benchmark['formats']['docx']['artifacts']['initial']['path'])
    if extension == 'docx':
        document.tables[0].cell(28, column).text = replacement
        mutated = tmp_path / 'mutated.docx'
        document.save(str(mutated))
        evidence = documents.inspect_docx(mutated)
    else:
        rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
        rows[28][column] = replacement
        tools = PDFToolSet(tmp_path)
        handle = documents.build_fixture(tools, 'mutated.pdf')
        tools.compose_artifact(handle, 'packages', {'kind': 'table', 'rows': rows})
        generated = tools.render_document(handle)
        evidence = documents.inspect_pdf(Path(generated.path))
    assert evidence['table_records_complete'] is False


def test_duplicate_record_rejects_integrated_document_gate(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    original = documents.build_fixture

    def duplicate_last_record(tools: Any, filename: str) -> Any:
        handle = original(tools, filename)
        rows = documents.expected_table_rows()
        rows[-1][0] = 'WP-01'
        tools.compose_artifact(handle, 'packages', {'kind': 'table', 'rows': rows})
        return handle

    monkeypatch.setattr(documents, 'build_fixture', duplicate_last_record)
    report = artifact_quality.run(tmp_path)
    assert report['automated']['status'] == 'fail'
    assert report['formats']['docx']['automated']['status'] == 'fail'
    assert report['formats']['pdf']['automated']['status'] == 'fail'
    assert report['formats']['pptx']['automated']['status'] == 'pass'
