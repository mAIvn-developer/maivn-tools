from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any, cast

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN

from evals.artifact_quality_presentations import run
from maivn_tools import PresentationsToolSet
from maivn_tools.connectors.presentations.models import PresentationHandle


def _table(tools: PresentationsToolSet, height: float) -> PresentationHandle:
    deck = tools.create_presentation('quality.pptx')
    tools.put_slide(deck, 'table', {'anchor': 'end'})
    ref = tools.compose_artifact(
        deck,
        'table',
        {
            'kind': 'table',
            'rows': [
                ['Owner', 'Decision gate'],
                ['Operations', 'Weekend coverage confirmed before regional launch'],
            ],
        },
    )
    tools.put_element(
        deck,
        'table',
        'table',
        {
            'type': 'table',
            'composition': ref,
            'x': 1,
            'y': 2,
            'width': 6,
            'height': height,
            'font_size': 18,
            'font_family': 'Arial',
            'alignment': 'right',
        },
    )
    return deck


def test_chart_axis_ids_open_in_powerpoint_signed_integer_range(tmp_path: Path) -> None:
    """Unsigned IDs above 2^31 were rejected by installed PowerPoint."""
    tools = PresentationsToolSet(tmp_path)
    deck = tools.create_presentation('chart.pptx')
    tools.put_slide(deck, 'metrics', {'anchor': 'end'})
    ref = tools.compose_artifact(
        deck,
        'chart',
        {
            'kind': 'chart',
            'categories': ['June', 'July'],
            'series': [{'name': 'Hours', 'values': [10.0, 7.0]}],
        },
    )
    tools.put_element(
        deck,
        'metrics',
        'chart',
        {
            'type': 'chart',
            'composition': ref,
            'x': 1,
            'y': 2,
            'width': 10,
            'height': 4,
        },
    )
    generated = tools.render_presentation(deck)
    with zipfile.ZipFile(generated.path) as archive:
        chart = archive.read('ppt/charts/chart1.xml')
    ids = [int(value) for value in re.findall(rb'<c:(?:axId|crossAx) val="(-?\d+)"', chart)]
    assert ids
    assert all(0 < value < (1 << 31) for value in ids)
    assert len(set(ids)) == 2
    assert ids[0] == ids[2] == ids[5]
    assert ids[1] == ids[3] == ids[4]


def test_table_respects_requested_font_and_alignment_for_every_paragraph(tmp_path: Path) -> None:
    tools = PresentationsToolSet(tmp_path)
    deck = _table(tools, 3.0)
    tools.compose_artifact(
        deck, 'table', {'kind': 'table', 'rows': [['Owner', 'Gate'], ['Ops', 'First\nSecond']]}
    )
    generated = tools.render_presentation(deck)
    table = cast('Any', Presentation(generated.path).slides[0].shapes[0]).table
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                assert paragraph.font.size.pt == 18
                assert paragraph.font.name == 'Arial'
                assert paragraph.alignment == PP_ALIGN.RIGHT


def test_layout_reports_crowded_table_cells_and_clears_after_resize(tmp_path: Path) -> None:
    tools = PresentationsToolSet(tmp_path)
    deck = _table(tools, 0.8)
    issues = tools.inspect_layout(deck)
    assert any(issue['code'] == 'table_text_overflow' for issue in issues)
    roomy = _table(tools, 3.0)
    assert not tools.inspect_layout(roomy)


@pytest.mark.parametrize(
    ('text', 'expected_code'),
    [
        ('Across the pilot (April\nJune), revenue was $37,000.', 'text_hard_break'),
        (r'April\u2013June revenue was \$37,000.', 'text_literal_escape'),
        ('April\u2013June revenue was $37,000.', None),
        ('Revenue increased.\nCosts stayed flat.', None),
    ],
)
def test_layout_flags_suspicious_prose_without_rewriting_it(
    tmp_path: Path, text: str, expected_code: str | None
) -> None:
    """Surface the benchmark text defects while preserving intentional author input."""
    tools = PresentationsToolSet(tmp_path)
    deck = tools.create_presentation('text.pptx')
    tools.put_slide(deck, 'conclusion', {'anchor': 'end'})
    ref = tools.compose_artifact(deck, 'body', {'kind': 'text', 'text': text})
    tools.put_element(
        deck,
        'conclusion',
        'body',
        {'type': 'text', 'composition': ref, 'x': 1, 'y': 1, 'width': 10, 'height': 4},
    )
    issues = tools.inspect_layout(deck)
    if expected_code is None:
        assert not issues
    else:
        assert {issue['code'] for issue in issues} == {expected_code}
        assert issues[0]['slide_id'] == 'conclusion'
        assert issues[0]['element_id'] == 'body'
    generated = tools.render_presentation(deck)
    rendered = cast('Any', Presentation(generated.path).slides[0].shapes[0]).text
    assert rendered.replace('\v', '\n') == text


def test_chart_legend_sits_below_plot_with_explicit_font(tmp_path: Path) -> None:
    result = run(tmp_path)
    deck = Presentation(str(result['initial_path']))
    chart = cast('Any', deck.slides[1].shapes[1]).chart
    assert chart.legend.position == XL_LEGEND_POSITION.BOTTOM
    assert chart.legend.include_in_layout is False
    assert chart.legend.font.name == 'Arial'
    assert chart.legend.font.size.pt == 18
    assert chart.category_axis.tick_labels.font.name == 'Arial'
    assert chart.value_axis.tick_labels.font.name == 'Arial'


def test_single_series_chart_displays_series_name_and_units(tmp_path: Path) -> None:
    tools = PresentationsToolSet(tmp_path)
    generated = tools.create_pptx(
        'single-series.pptx',
        {
            'slides': [
                {
                    'slide_id': 'metrics',
                    'elements': [
                        {
                            'element_id': 'revenue',
                            'type': 'chart',
                            'content': {
                                'kind': 'chart',
                                'categories': ['January', 'February'],
                                'series': [{'name': 'Revenue (USD)', 'values': [100.0, 120.0]}],
                            },
                            'x': 1,
                            'y': 2,
                            'width': 10,
                            'height': 4,
                        }
                    ],
                }
            ]
        },
    )
    chart = cast('Any', Presentation(generated.path).slides[0].shapes[0]).chart
    assert chart.has_legend is True
    assert chart.series[0].name == 'Revenue (USD)'
    assert chart.plots[0].vary_by_categories is False
    assert chart.legend.position == XL_LEGEND_POSITION.BOTTOM
    assert chart.legend.include_in_layout is False
