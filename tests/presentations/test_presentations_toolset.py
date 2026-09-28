# pyright: strict
from __future__ import annotations

import zipfile
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from typing import Any, cast

import pytest
from maivn import Agent
from maivn._internal.compat.decorators import (
    ARGUMENT_PRODUCER_REQUIREMENTS_ATTR,
)
from pptx import Presentation

import maivn_tools
from maivn_tools.connectors.presentations.models import PresentationHandle


def test_presentation_tools_publish_state_and_composition_dependency_contracts(
    tmp_path: Path,
) -> None:
    tools = Agent(name='presentations-contract', api_key='test-key').add_toolset(
        maivn_tools.PresentationsToolSet(tmp_path)
    )
    by_name = {tool.name: tool for tool in tools}

    read_schema = by_name['PRESENTATIONS_read_presentation'].output_schema
    assert isinstance(read_schema, dict)
    assert read_schema['required'] == ['presentation', 'compositions', 'slides']
    assert not getattr(
        by_name['PRESENTATIONS_put_element'].target,
        ARGUMENT_PRODUCER_REQUIREMENTS_ATTR,
        (),
    )


def _build_presentation(
    toolset: maivn_tools.PresentationsToolSet, filename: str
) -> PresentationHandle:
    deck = toolset.create_presentation(filename)
    toolset.put_slide(deck, 'intro', {'anchor': 'end'}, layout='blank')
    title = toolset.compose_artifact(deck, 'title', {'kind': 'text', 'text': 'Artifact Toolsets'})
    bullets = toolset.compose_artifact(
        deck, 'bullets', {'kind': 'list', 'items': ['Compose', 'Inspect', 'Render']}
    )
    chart = toolset.compose_artifact(
        deck,
        'chart',
        {
            'kind': 'chart',
            'categories': ['Compose', 'Inspect', 'Render'],
            'series': [{'name': 'Confidence', 'values': [70.0, 85.0, 100.0]}],
        },
    )
    toolset.put_element(
        deck,
        'intro',
        'title',
        {'type': 'title', 'composition': title, 'x': 0.7, 'y': 0.5, 'width': 11.9, 'height': 0.8},
    )
    toolset.put_element(
        deck,
        'intro',
        'chart',
        {'type': 'chart', 'composition': chart, 'x': 7.0, 'y': 1.7, 'width': 5.4, 'height': 3.5},
    )
    toolset.put_element(
        deck,
        'intro',
        'bullets',
        {'type': 'list', 'composition': bullets, 'x': 0.9, 'y': 1.7, 'width': 5.8, 'height': 3.5},
    )
    return deck


def test_presentation_chain_reopens_and_is_byte_deterministic(tmp_path: Path) -> None:
    left = maivn_tools.PresentationsToolSet(tmp_path / 'left')
    right = maivn_tools.PresentationsToolSet(tmp_path / 'right')
    left_deck = _build_presentation(left, 'proof.pptx')
    right_deck = _build_presentation(right, 'proof.pptx')

    left_file = left.render_presentation(left_deck)
    right_file = right.render_presentation(right_deck)
    assert left_file.sha256 == right_file.sha256
    deck = Presentation(left_file.path)
    assert len(deck.slides) == 1
    assert 'Artifact Toolsets' in '\n'.join(
        cast('str', cast('Any', shape).text)
        for shape in deck.slides[0].shapes
        if bool(cast('Any', shape).has_text_frame)
    )
    assert any(bool(cast('Any', shape).has_chart) for shape in deck.slides[0].shapes)
    with zipfile.ZipFile(left_file.path) as archive:
        assert archive.namelist() == sorted(archive.namelist())
        assert {entry.date_time for entry in archive.infolist()} == {(1980, 1, 1, 0, 0, 0)}
        chart_xml = archive.read('ppt/charts/chart1.xml')
        assert b'<c:axId val="-' not in chart_xml
        assert b'<c:crossAx val="-' not in chart_xml
        with zipfile.ZipFile(
            BytesIO(archive.read('ppt/embeddings/Microsoft_Excel_Sheet1.xlsx'))
        ) as chart_workbook:
            core_properties = chart_workbook.read('docProps/core.xml')
            assert b'>2000-01-01T00:00:00Z</dcterms:created>' in core_properties
            assert b'>2000-01-01T00:00:00Z</dcterms:modified>' in core_properties


def test_presentation_upsert_remove_readback_and_inline_refusal(tmp_path: Path) -> None:
    decks = maivn_tools.PresentationsToolSet(tmp_path)
    deck = decks.create_presentation('proof.pptx')
    decks.put_slide(deck, 'slide', {'anchor': 'end'}, layout='blank')
    first = decks.compose_artifact(deck, 'body', {'kind': 'text', 'text': 'First'})
    revised = decks.compose_artifact(deck, 'body', {'kind': 'text', 'text': 'Revised'})
    assert first == revised
    decks.put_element(
        deck,
        'slide',
        'body',
        {'type': 'text', 'composition': revised, 'x': 1.0, 'y': 1.0, 'width': 4.0, 'height': 1.0},
    )
    state = cast('dict[str, Any]', decks.read_presentation(deck))
    assert state['compositions'][0]['content']['text'] == 'Revised'
    assert state['slides'][0]['elements'][0]['element_id'] == 'body'
    assert decks.remove_element(deck, 'slide', 'body')['removed'] is True
    with pytest.raises(
        maivn_tools.CompositionDependencyError, match='PRESENTATIONS_compose_artifact'
    ):
        decks.put_element(
            deck,
            'slide',
            'inline',
            {'type': 'text', 'composition': 'raw', 'x': 1.0, 'y': 1.0, 'width': 4.0, 'height': 1.0},  # type: ignore[typeddict-item]
        )


def test_presentation_one_shot_routes_through_compositions(tmp_path: Path) -> None:
    decks = maivn_tools.PresentationsToolSet(tmp_path)
    generated = decks.create_pptx(
        'small.pptx',
        {
            'slides': [
                {
                    'slide_id': 'one',
                    'elements': [
                        {
                            'element_id': 'title',
                            'type': 'title',
                            'content': {'kind': 'text', 'text': 'One-shot'},
                            'x': 0.7,
                            'y': 0.5,
                            'width': 11.9,
                            'height': 0.8,
                        }
                    ],
                }
            ]
        },
    )
    state = cast(
        'dict[str, Any]',
        decks.read_presentation(cast('PresentationHandle', generated.workspace)),
    )
    assert state['compositions'][0]['composition_id'] == 'one-title-content'


def test_presentation_handles_are_isolated_under_concurrency(tmp_path: Path) -> None:
    decks = maivn_tools.PresentationsToolSet(tmp_path)

    def create_handle(_: int) -> PresentationHandle:
        return decks.create_presentation('same.pptx')

    with ThreadPoolExecutor(max_workers=4) as executor:
        handles = list(executor.map(create_handle, range(8)))

    assert len({handle['presentation_id'] for handle in handles}) == 8
