from pathlib import Path
from typing import cast

from docx import Document
from docx.shared import Pt
from docx.styles.style import ParagraphStyle

from evals.artifact_quality_documents import build_fixture, inspect_docx
from maivn_tools import DocumentsToolSet


def test_report_tables_repeat_headers_keep_rows_and_allocate_prose_width(tmp_path: Path) -> None:
    tools = DocumentsToolSet(tmp_path)
    handle = build_fixture(tools, 'report.docx')
    result = tools.render_document(handle)
    document = Document(result.path)
    table = document.tables[0]
    assert inspect_docx(Path(result.path))['table_headers_repeat']
    prose_width, key_width = table.columns[1].width, table.columns[0].width
    assert prose_width is not None and key_width is not None
    assert prose_width > key_width
    assert len(document.element.xpath('./w:body/w:tbl/w:tr/w:trPr/w:cantSplit')) == len(table.rows)
    assert table.cell(0, 0).paragraphs[0].runs[0].bold
    assert document.element.xpath('./w:body/w:tbl/w:tblPr/w:tblBorders/w:insideH')


def test_document_defaults_set_readable_typography_and_unadorned_headings(tmp_path: Path) -> None:
    tools = DocumentsToolSet(tmp_path)
    handle = build_fixture(tools, 'report.docx')
    document = Document(tools.render_document(handle).path)
    normal = cast('ParagraphStyle', document.styles['Normal'])
    assert normal.font.name == 'Arial'
    assert normal.font.size == Pt(11)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        style = cast('ParagraphStyle', document.styles[name])
        assert str(style.font.color.rgb) == '000000'
        assert style.element.xpath('./w:pPr/w:keepNext')
        assert not style.element.xpath('./w:pPr/w:pBdr')
