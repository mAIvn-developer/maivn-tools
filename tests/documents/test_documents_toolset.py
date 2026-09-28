# pyright: strict
from __future__ import annotations

import base64
import hashlib
import json
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import cast

import pytest
from docx import Document
from maivn import Agent
from maivn import toolset_options as get_toolset_options
from maivn._internal.compat.decorators import (
    ARGUMENT_PRODUCER_REQUIREMENTS_ATTR,
)
from pydantic import TypeAdapter, ValidationError

import maivn_tools
from maivn_tools.connectors.documents import toolset as documents_module
from maivn_tools.connectors.documents.models import (
    CompositionContent,
    DocumentBlock,
    DocumentHandle,
)


def test_document_workspace_refuses_use_after_crossing_a_process_boundary(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    doc = documents.create_document('report.docx')
    owner_pid = documents_module.os.getpid()
    monkeypatch.setattr(documents_module.os, 'getpid', lambda: owner_pid + 1)

    with pytest.raises(RuntimeError, match='single-process-per-workspace'):
        documents.read_document(doc)


def _compose_rich_document(
    documents: maivn_tools.DocumentsToolSet,
    filename: str,
) -> DocumentHandle:
    doc = documents.create_document(filename)
    authored: list[tuple[str, CompositionContent]] = [
        ('title', {'kind': 'text', 'text': 'Artifact Toolsets'}),
        ('heading', {'kind': 'text', 'text': 'Working shape'}),
        ('paragraph', {'kind': 'text', 'text': 'The model plans; the toolset builds.'}),
        ('bullets', {'kind': 'list', 'items': ['Persist state', 'Retry safely']}),
        ('numbers', {'kind': 'list', 'items': ['Compose', 'Place', 'Render']}),
        ('table', {'kind': 'table', 'rows': [['Gate', 'Result'], ['DOCX', 'real file']]}),
    ]
    references = {
        composition_id: documents.compose_artifact(doc, composition_id, content)
        for composition_id, content in authored
    }
    blocks: list[tuple[str, DocumentBlock]] = [
        ('title', {'type': 'title', 'composition': references['title']}),
        (
            'heading',
            {'type': 'heading', 'level': 1, 'composition': references['heading']},
        ),
        ('paragraph', {'type': 'paragraph', 'composition': references['paragraph']}),
        ('bullets', {'type': 'bullet_list', 'composition': references['bullets']}),
        ('numbers', {'type': 'numbered_list', 'composition': references['numbers']}),
        ('table', {'type': 'table', 'composition': references['table']}),
        ('page-break', {'type': 'page_break'}),
    ]
    for block_id, block in blocks:
        documents.put_block(doc, block_id, block, {'anchor': 'end'})
    return doc


def test_documents_toolset_is_a_public_toolset() -> None:
    assert hasattr(maivn_tools, 'DocumentsToolSet')

    options = get_toolset_options(maivn_tools.DocumentsToolSet)
    assert options is not None
    assert options.prefix == 'documents'
    assert maivn_tools.DocumentsToolSet.metadata.name == 'documents'
    assert maivn_tools.DocumentsToolSet.metadata.display_name == 'Documents'


def test_create_document_returns_distinct_persisted_handles(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)

    first = documents.create_document('report.docx')
    second = documents.create_document('report.docx')

    assert first['document_id'] != second['document_id']
    assert first['filename'] == 'report.docx'
    assert second['filename'] == 'report.docx'


def test_compose_artifact_upserts_by_caller_chosen_id_and_reads_back(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    doc = documents.create_document('report.docx')

    first = documents.compose_artifact(
        doc,
        'intro',
        {'kind': 'text', 'text': 'First draft'},
    )
    retried = documents.compose_artifact(
        doc,
        'intro',
        {'kind': 'text', 'text': 'First draft'},
    )
    revised = documents.compose_artifact(
        doc,
        'intro',
        {'kind': 'text', 'text': 'Revised draft'},
    )

    assert first == retried == revised
    assert documents.read_document(doc) == {
        'document': doc,
        'compositions': [
            {
                'composition_id': 'intro',
                'content': {'kind': 'text', 'text': 'Revised draft'},
            }
        ],
        'outline': [],
    }


def test_compose_artifact_validates_ids_and_structured_content(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    doc = documents.create_document('report.docx')

    with pytest.raises(ValueError, match='composition_id'):
        documents.compose_artifact(doc, '', {'kind': 'text', 'text': 'content'})
    with pytest.raises(ValueError, match='same number of columns'):
        documents.compose_artifact(
            doc,
            'ragged',
            {'kind': 'table', 'rows': [['one', 'two'], ['only-one']]},
        )


def test_put_block_refuses_inline_content_with_honest_dependency_error(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    doc = documents.create_document('report.docx')

    with pytest.raises(
        maivn_tools.CompositionDependencyError,
        match=r"argument 'block\.composition'.*DOCUMENTS_compose_artifact",
    ):
        documents.put_block(
            doc,
            'intro',
            cast('DocumentBlock', {'type': 'paragraph', 'composition': 'inline prose'}),
            {'anchor': 'end'},
        )

    with pytest.raises(
        maivn_tools.CompositionDependencyError,
        match=r"argument 'block\.composition'.*DOCUMENTS_compose_artifact",
    ):
        documents.put_block(
            doc,
            'missing-content',
            cast('DocumentBlock', {'type': 'paragraph'}),
            {'anchor': 'end'},
        )


def test_put_block_uses_persisted_composition_authority_without_current_turn_gate(
    tmp_path: Path,
) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    agent = Agent(name='documents-contract', api_key='test-key')

    tools = agent.add_toolset(documents)
    put_block = next(tool for tool in tools if tool.name == 'DOCUMENTS_put_block')

    assert not getattr(put_block.target, ARGUMENT_PRODUCER_REQUIREMENTS_ATTR, ())
    assert 'composition_dependencies' not in put_block.metadata
    block_schema = put_block.input_schema['properties']['block']
    assert isinstance(block_schema, dict)
    names = [
        'DocumentContentBlock',
        'DocumentHeadingBlock',
        'DocumentImageBlock',
        'DocumentPageBreak',
    ]
    assert block_schema['anyOf'] == [{'$ref': f'#/$defs/{name}'} for name in names]
    definitions = put_block.input_schema['$defs']
    assert definitions['DocumentHeadingBlock']['required'] == ['type', 'composition', 'level']
    assert definitions['DocumentHeadingBlock']['properties']['level']['enum'] == list(range(1, 10))
    assert definitions['DocumentImageBlock']['required'] == [
        'type',
        'composition',
        'image_placement',
    ]


def test_document_tools_publish_first_class_artifact_output_schemas(tmp_path: Path) -> None:
    agent = Agent(name='documents-output-contract', api_key='test-key')
    tools = agent.add_toolset(maivn_tools.DocumentsToolSet(tmp_path))
    schemas = {tool.name: tool.output_schema for tool in tools}

    compose_schema = schemas['DOCUMENTS_compose_artifact']
    render_schema = schemas['DOCUMENTS_render_document']
    assert isinstance(compose_schema, dict)
    assert compose_schema['required'] == ['document_id', 'composition_id']
    assert isinstance(render_schema, dict)
    assert {
        'kind',
        'artifact_kind',
        'filename',
        'path',
        'mime_type',
        'size_bytes',
        'sha256',
        'custody',
    } <= set(cast('list[object]', render_schema['required']))
    assert render_schema['properties']['kind']['const'] == 'generated_file'
    assert 'workspace' in render_schema['properties']
    assert 'content_base64' in render_schema['properties']


def test_put_block_retries_converge_and_revisions_preserve_position(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    doc = documents.create_document('report.docx')
    intro = documents.compose_artifact(doc, 'intro-v1', {'kind': 'text', 'text': 'Intro'})
    closing = documents.compose_artifact(doc, 'closing', {'kind': 'text', 'text': 'Close'})

    documents.put_block(
        doc,
        'intro',
        {'type': 'paragraph', 'composition': intro},
        {'anchor': 'end'},
    )
    documents.put_block(
        doc,
        'closing',
        {'type': 'paragraph', 'composition': closing},
        {'anchor': 'end'},
    )
    documents.put_block(
        doc,
        'intro',
        {'type': 'paragraph', 'composition': intro},
        {'anchor': 'end'},
    )

    assert [entry['block_id'] for entry in documents.read_document(doc)['outline']] == [
        'intro',
        'closing',
    ]

    revised = documents.compose_artifact(
        doc,
        'intro-v2',
        {'kind': 'text', 'text': 'Revised intro'},
    )
    documents.put_block(
        doc,
        'intro',
        {'type': 'paragraph', 'composition': revised},
        {'anchor': 'end'},
    )
    state = documents.read_document(doc)
    assert [entry['block_id'] for entry in state['outline']] == ['intro', 'closing']
    assert state['outline'][0]['block'] == {
        'type': 'paragraph',
        'composition': revised,
    }

    documents.put_block(
        doc,
        'intro',
        {'type': 'paragraph', 'composition': revised},
        {'anchor': 'after', 'block_id': 'closing'},
    )
    assert [entry['block_id'] for entry in documents.read_document(doc)['outline']] == [
        'closing',
        'intro',
    ]


def test_remove_block_is_retry_safe_and_page_break_needs_no_composition(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    doc = documents.create_document('report.docx')
    documents.put_block(doc, 'break', {'type': 'page_break'}, {'anchor': 'start'})

    assert documents.remove_block(doc, 'break')['removed'] is True
    assert documents.remove_block(doc, 'break')['removed'] is False
    assert documents.read_document(doc)['outline'] == []


def test_render_document_builds_a_real_hash_verified_docx(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    doc = _compose_rich_document(documents, 'proof.docx')

    generated = documents.render_document(doc, include_bytes=True)
    generated_path = Path(generated.path)
    raw = generated_path.read_bytes()

    assert generated_path.is_absolute()
    assert generated_path.is_relative_to(tmp_path.resolve())
    assert generated.kind == 'generated_file'
    assert generated.artifact_kind == 'document'
    assert generated.filename == 'proof.docx'
    assert generated.workspace is not None
    assert generated_path.parent.name == generated.workspace['source_revision']
    assert {key: generated.workspace[key] for key in doc} == doc
    assert len(generated.workspace['source_revision']) == 64
    assert generated.custody == 'ordinary'
    assert generated.mime_type == (
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    assert generated.size_bytes == len(raw) > 1_000
    assert generated.sha256 == hashlib.sha256(raw).hexdigest()
    content_base64 = generated.content_base64
    assert isinstance(content_base64, str)
    assert base64.b64decode(content_base64) == raw
    assert zipfile.is_zipfile(generated_path)

    opened = Document(str(generated_path))
    paragraph_text = [paragraph.text for paragraph in opened.paragraphs]
    assert paragraph_text[:8] == [
        'Artifact Toolsets',
        'Working shape',
        'The model plans; the toolset builds.',
        'Persist state',
        'Retry safely',
        'Compose',
        'Place',
        'Render',
    ]
    assert opened.tables[0].cell(1, 1).text == 'real file'

    with zipfile.ZipFile(generated_path) as archive:
        assert archive.namelist() == sorted(archive.namelist())
        assert {entry.date_time for entry in archive.infolist()} == {(1980, 1, 1, 0, 0, 0)}


def test_identical_build_sequences_produce_byte_identical_docx_files(tmp_path: Path) -> None:
    left_dir = tmp_path / 'left'
    right_dir = tmp_path / 'right'
    left_dir.mkdir()
    right_dir.mkdir()
    left = maivn_tools.DocumentsToolSet(left_dir)
    right = maivn_tools.DocumentsToolSet(right_dir)

    left_file = left.render_document(_compose_rich_document(left, 'same.docx'))
    right_file = right.render_document(_compose_rich_document(right, 'same.docx'))

    assert left_file.sha256 == right_file.sha256
    assert Path(left_file.path).read_bytes() == Path(right_file.path).read_bytes()


def test_concurrent_handles_and_disjoint_block_operations_do_not_collide(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)

    def create_handle(index: int) -> DocumentHandle:
        doc = documents.create_document('shared-name.docx')
        documents.compose_artifact(
            doc,
            'body',
            {'kind': 'text', 'text': f'Document {index}'},
        )
        return doc

    with ThreadPoolExecutor(max_workers=8) as executor:
        handles = list(executor.map(create_handle, range(16)))

    assert len({handle['document_id'] for handle in handles}) == 16
    authored_texts = [
        documents.read_document(handle)['compositions'][0]['content'].get('text')
        for handle in handles
    ]
    assert all(isinstance(text, str) for text in authored_texts)
    assert set(authored_texts) == {f'Document {index}' for index in range(16)}

    shared = documents.create_document('dag.docx')
    references = [
        documents.compose_artifact(
            shared,
            f'composition-{index}',
            {'kind': 'text', 'text': f'Block {index}'},
        )
        for index in range(16)
    ]

    def put_disjoint(index: int) -> None:
        separate_instance = maivn_tools.DocumentsToolSet(tmp_path)
        separate_instance.put_block(
            shared,
            f'block-{index}',
            {'type': 'paragraph', 'composition': references[index]},
            {'anchor': 'end'},
        )

    with ThreadPoolExecutor(max_workers=8) as executor:
        list(executor.map(put_disjoint, range(16)))

    assert {entry['block_id'] for entry in documents.read_document(shared)['outline']} == {
        f'block-{index}' for index in range(16)
    }


def test_create_document_rejects_unsafe_or_non_docx_filenames(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)

    with pytest.raises(ValueError, match='directly beneath output_dir'):
        documents.create_document('../escape.docx')
    with pytest.raises(ValueError, match=r'\.docx'):
        documents.create_document('not-a-document.txt')


def test_document_handle_cannot_escape_the_persisted_workspace(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    (tmp_path / 'escape.json').write_text(
        json.dumps(
            {
                'version': 1,
                'document_id': '../escape',
                'filename': 'proof.docx',
                'compositions': {},
                'blocks': [],
            }
        ),
        encoding='utf-8',
    )

    forged = cast('DocumentHandle', {'document_id': '../escape', 'filename': 'proof.docx'})
    with pytest.raises(ValueError, match='invalid document handle'):
        documents.read_document(forged)


def test_render_document_refuses_overwrite_without_explicit_permission(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    first = documents.create_document('same.docx')
    second = documents.create_document('same.docx')
    documents.render_document(first)

    with pytest.raises(FileExistsError):
        documents.render_document(second)

    replaced = documents.render_document(second, overwrite=True)
    assert Path(replaced.path).is_file()


def test_render_document_refuses_a_composition_revised_to_the_wrong_kind(
    tmp_path: Path,
) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    doc = documents.create_document('invalid.docx')
    reference = documents.compose_artifact(doc, 'body', {'kind': 'text', 'text': 'Body'})
    documents.put_block(
        doc,
        'body',
        {'type': 'paragraph', 'composition': reference},
        {'anchor': 'end'},
    )
    documents.compose_artifact(doc, 'body', {'kind': 'list', 'items': ['Changed kind']})

    with pytest.raises(ValueError, match=r'paragraph.*requires a text composition'):
        documents.render_document(doc)


def test_create_docx_routes_inline_content_through_persisted_compositions(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)

    generated = documents.create_docx(
        'small.docx',
        {
            'blocks': [
                {'block_id': 'title', 'type': 'title', 'content': 'Small document'},
                {
                    'block_id': 'items',
                    'type': 'bullet_list',
                    'items': ['One', 'Two'],
                },
                {'block_id': 'break', 'type': 'page_break'},
            ]
        },
    )

    assert generated.workspace is not None
    state = documents.read_document(cast('DocumentHandle', generated.workspace))
    assert [item['composition_id'] for item in state['compositions']] == [
        'items-content',
        'title-content',
    ]
    assert [entry['block_id'] for entry in state['outline']] == ['title', 'items', 'break']
    assert Path(generated.path).is_file()


def test_create_docx_failure_does_not_leave_an_unreachable_workspace(tmp_path: Path) -> None:
    documents = maivn_tools.DocumentsToolSet(tmp_path)
    existing = documents.create_document('same.docx')
    documents.render_document(existing)
    workspace_dir = tmp_path / '.maivn-documents'
    before = {path.name for path in workspace_dir.glob('*.json')}

    with pytest.raises(FileExistsError):
        documents.create_docx('same.docx', {'blocks': []})

    assert {path.name for path in workspace_dir.glob('*.json')} == before


@pytest.mark.parametrize(
    'block',
    [
        {'type': 'paragraph'},
        {'type': 'heading'},
        {'type': 'heading', 'composition': {'document_id': 'doc', 'composition_id': 'c'}},
        {
            'type': 'heading',
            'composition': {'document_id': 'doc', 'composition_id': 'c'},
            'level': 0,
        },
        {
            'type': 'heading',
            'composition': {'document_id': 'doc', 'composition_id': 'c'},
            'level': 10,
        },
        {'type': 'image', 'composition': {'document_id': 'doc', 'composition_id': 'c'}},
    ],
)
def test_document_placement_schema_rejects_runtime_invalid_shapes(block: dict[str, object]) -> None:
    adapter: TypeAdapter[DocumentBlock] = TypeAdapter(DocumentBlock)
    with pytest.raises(ValidationError):
        adapter.validate_python(block)


def test_document_placement_schema_accepts_heading_and_page_break() -> None:
    adapter: TypeAdapter[DocumentBlock] = TypeAdapter(DocumentBlock)
    assert adapter.validate_python({'type': 'page_break'}) == {'type': 'page_break'}
    heading = {
        'type': 'heading',
        'level': 1,
        'composition': {'document_id': 'doc', 'composition_id': 'c'},
    }
    assert adapter.validate_python(heading) == heading
