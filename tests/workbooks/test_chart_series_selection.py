"""Chart selection preserves categories without plotting unrelated measures."""

from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter

from maivn_tools import WorkbooksToolSet
from maivn_tools.connectors.workbooks.models import RangeFormat


@pytest.mark.parametrize('columns', [None, [4], [4, 2]])
def test_chart_uses_selected_sheet_columns_in_declared_order(
    tmp_path: Path, columns: list[int] | None
) -> None:
    tools = WorkbooksToolSet(tmp_path)
    book = tools.create_workbook('measures.xlsx')
    tools.put_sheet(book, 'data', 'Data', {'anchor': 'end'})
    data = tools.compose_artifact(
        book,
        'table',
        {'kind': 'matrix', 'values': [['Month', 'Users', 'ARPU', 'Revenue'], ['Jan', 20, 5, 100]]},
    )
    chart: dict[str, object] = {'type': 'line', 'title': 'Revenue', 'anchor': 'F2'}
    if columns is not None:
        chart['series_columns'] = columns
    tools.put_range(book, 'table', 'data', 'A1', data, cast('RangeFormat', {'chart': chart}))
    result = tools.render_workbook(book)
    sheet = load_workbook(result.path)['Data']
    [rendered] = cast('Any', sheet)._charts  # noqa: SLF001 - openpyxl has no public chart reader.
    wanted = columns or [2, 3, 4]
    assert len(rendered.series) == len(wanted)
    for series, column in zip(rendered.series, wanted, strict=True):
        letter = get_column_letter(column)
        assert series.tx.strRef.f == f"'Data'!{letter}1"
        assert series.val.numRef.f == f"'Data'!${letter}$2"
        assert series.cat.strRef.f == "'Data'!$A$2"
        assert series.cat.strRef.strCache.pt[0].v == 'Jan'
    assert (rendered.legend is None) == (len(wanted) == 1)


@pytest.mark.parametrize('columns', [[], [0], [16385], [1], [5], [2, 2]])
def test_invalid_chart_series_does_not_change_workbook(tmp_path: Path, columns: list[int]) -> None:
    tools = WorkbooksToolSet(tmp_path)
    book = tools.create_workbook('invalid-series.xlsx')
    tools.put_sheet(book, 'data', 'Data', {'anchor': 'end'})
    ref = tools.compose_artifact(book, 'header', {'kind': 'row', 'values': ['Month', 'Revenue']})
    before = tools.read_workbook(book)
    formatting = cast(
        'RangeFormat',
        {
            'chart': {
                'type': 'line',
                'title': 'Revenue',
                'anchor': 'F2',
                'source_range': 'A1:D4',
                'series_columns': columns,
            }
        },
    )
    with pytest.raises(ValueError):
        tools.put_range(book, 'header', 'data', 'A1', ref, formatting)
    assert tools.read_workbook(book) == before


def test_resize_cannot_remove_a_selected_chart_column(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    book = tools.create_workbook('resize.xlsx')
    tools.put_sheet(book, 'data', 'Data', {'anchor': 'end'})
    ref = tools.compose_artifact(
        book, 'data', {'kind': 'matrix', 'values': [['Month', 'Count', 'Revenue'], ['Jan', 1, 2]]}
    )
    tools.put_range(
        book,
        'data',
        'data',
        'A1',
        ref,
        {
            'chart': {'type': 'line', 'title': 'Revenue', 'anchor': 'F2', 'series_columns': [3]},
        },
    )
    before = tools.read_workbook(book)
    with pytest.raises(ValueError, match='series columns'):
        tools.compose_artifact(
            book, 'data', {'kind': 'matrix', 'values': [['Month', 'Count'], ['Jan', 1]]}
        )
    assert tools.read_workbook(book) == before


def test_readback_reports_existing_destination_without_overwriting_it(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    book = tools.create_workbook('shared-name.xlsx')
    tools.put_sheet(book, 'data', 'Data', {'anchor': 'end'})
    ref = tools.compose_artifact(book, 'data', {'kind': 'scalar', 'value': 3})
    tools.put_range(book, 'data', 'data', 'A1', ref, {})
    assert tools.read_workbook(book).get('output_file_exists') is False
    target = tmp_path / 'shared-name.xlsx'
    target.write_bytes(b'Existing developer output')
    assert tools.read_workbook(book).get('output_file_exists') is True
    with pytest.raises(FileExistsError):
        tools.render_workbook(book)
    assert target.read_bytes() == b'Existing developer output'
