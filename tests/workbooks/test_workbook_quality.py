"""Quality regressions exercised through the public workbook toolset."""

from pathlib import Path
from typing import Any, cast
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pydantic import TypeAdapter

from maivn_tools import WorkbooksToolSet
from maivn_tools.connectors.workbooks.models import RangeFormat


@pytest.mark.parametrize(
    'formula',
    [
        '=Assumptions!$B$2*B2',
        "='Growth assumptions'!$B$3*(1+B2)",
        '=SUM(B2:B13)/COUNT(B2:B13)',
        '=B2*(1+15%)',
        '=SUM($B$2:$B$13,MAX(C2:C13))',
    ],
)
def test_forecast_formulas_can_use_internal_assumptions(tmp_path: Path, formula: str) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('forecast.xlsx')
    tools.put_sheet(handle, 'forecast', 'Forecast', {'anchor': 'end'})
    ref = tools.compose_artifact(handle, 'forecast', {'kind': 'formulas', 'formulas': [[formula]]})
    tools.put_range(handle, 'forecast', 'forecast', 'B2', ref, {})
    result = tools.render_workbook(handle)
    book = load_workbook(result.path)
    assert book['Forecast']['B2'].value == formula


@pytest.mark.parametrize(
    'formula',
    [
        '=WEBSERVICE("https://example.test")',
        "='[external.xlsx]Sheet'!A1",
        '=SUM(A1:) ',
        '=SUM(A1,,A2)',
        '=A1+',
        '=A0+A2',
        '=XFE1+A1',
        '=A1048577+A1',
        '=SUM(A1:A2))',
        '=cmd|stuff!A1',
        '=INDIRECT("A1")',
    ],
)
def test_forecast_formulas_refuse_invalid_or_external_expressions(
    tmp_path: Path,
    formula: str,
) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('forecast.xlsx')
    with pytest.raises(ValueError, match='formula'):
        tools.compose_artifact(handle, 'bad', {'kind': 'formulas', 'formulas': [[formula]]})


def test_data_strings_do_not_become_unvalidated_formulas(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('data.xlsx')
    tools.put_sheet(handle, 'data', 'Data', {'anchor': 'end'})
    literal = '=this is a literal customer label'
    ref = tools.compose_artifact(handle, 'literal', {'kind': 'row', 'values': [literal]})
    tools.put_range(handle, 'literal', 'data', 'A1', ref, {})
    result = tools.render_workbook(handle)
    cell = load_workbook(result.path)['Data']['A1']
    assert cell.value == literal
    assert cell.data_type == 's'


def test_column_width_keeps_long_heading_after_short_later_range(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('readable.xlsx')
    tools.put_sheet(handle, 'data', 'Data', {'anchor': 'end'})
    heading = tools.compose_artifact(
        handle,
        'heading',
        {'kind': 'row', 'values': ['Monthly subscription revenue']},
    )
    value = tools.compose_artifact(handle, 'value', {'kind': 'scalar', 'value': 0})
    tools.put_range(handle, 'heading', 'data', 'A1', heading, {'style': 'header_row'})
    tools.put_range(handle, 'value', 'data', 'A2', value, {'number_format': '$#,##0.00'})
    result = tools.render_workbook(handle)
    sheet = load_workbook(result.path)['Data']
    assert sheet.column_dimensions['A'].width >= 29
    assert cast('Any', sheet.row_dimensions[1]).height is not None
    assert sheet.sheet_view.showGridLines is False


def test_long_note_row_expands_and_survives_revision(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('notes.xlsx')
    tools.put_sheet(handle, 'notes', 'Notes', {'anchor': 'end'}, freeze_panes='A2')
    note = 'This forecast is illustrative. ' * 6
    ref = tools.compose_artifact(handle, 'note', {'kind': 'row', 'values': [note]})
    tools.put_range(handle, 'note', 'notes', 'A2', ref, {})
    initial = tools.render_workbook(handle)
    before = Path(initial.path).read_bytes()
    reopened = WorkbooksToolSet(tmp_path)
    reopened.compose_artifact(handle, 'note', {'kind': 'row', 'values': [note + 'Updated.']})
    revised = reopened.render_workbook(handle, overwrite=True)
    sheet = load_workbook(revised.path)['Notes']
    assert cast('Any', sheet.row_dimensions[2]).height is not None
    assert cast('Any', sheet.row_dimensions[2]).height >= 75
    assert sheet.freeze_panes == 'A2'
    assert Path(initial.path).read_bytes() == before


def test_chart_can_reference_separately_formatted_formula_cells(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('chart.xlsx')
    tools.put_sheet(handle, 'forecast', 'Forecast', {'anchor': 'end'})
    header = tools.compose_artifact(
        handle, 'header', {'kind': 'row', 'values': ['Month', 'Revenue']}
    )
    tools.put_range(
        handle,
        'header',
        'forecast',
        'A1',
        header,
        {
            'style': 'header_row',
            'chart': {
                'type': 'line',
                'title': 'Revenue (USD)',
                'anchor': 'D2',
                'source_range': 'A1:B3',
            },
        },
    )
    labels = tools.compose_artifact(
        handle, 'labels', {'kind': 'matrix', 'values': [['Jan'], ['Feb']]}
    )
    formulas = tools.compose_artifact(
        handle, 'values', {'kind': 'formulas', 'formulas': [['=100*2'], ['=100*3']]}
    )
    tools.put_range(handle, 'labels', 'forecast', 'A2', labels, {})
    tools.put_range(handle, 'values', 'forecast', 'B2', formulas, {'number_format': '$#,##0.00'})
    result = tools.render_workbook(handle)
    with ZipFile(result.path) as archive:
        chart = archive.read('xl/charts/chart1.xml')
    assert b"'Forecast'!$B$2:$B$3" in chart
    assert b"'Forecast'!$A$2:$A$3" in chart
    root = ET.fromstring(chart)  # noqa: S314 - XML was generated by the local renderer.
    ns = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    assert len(root.findall('.//c:delete[@val="0"]', ns)) == 2
    assert root.find('.//c:tickLblPos[@val="nextTo"]', ns) is not None


def test_expanding_composition_sizes_new_rows(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('expanding.xlsx')
    tools.put_sheet(handle, 'data', 'Data', {'anchor': 'end'})
    ref = tools.compose_artifact(handle, 'note', {'kind': 'matrix', 'values': [['Short']]})
    tools.put_range(handle, 'note', 'data', 'A1', ref, {})
    tools.compose_artifact(
        handle,
        'note',
        {
            'kind': 'matrix',
            'values': [
                ['Short'],
                ['A note with detailed operating assumptions. ' * 8],
            ],
        },
    )
    result = tools.render_workbook(handle)
    sheet = load_workbook(result.path)['Data']
    assert sheet.column_dimensions['A'].width == 40
    assert cast('Any', sheet.row_dimensions[2]).height >= 100


def test_mixed_value_and_formula_content_is_not_silently_discarded(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('mixed.xlsx')
    with pytest.raises(ValueError, match='separate'):
        tools.compose_artifact(
            handle,
            'mixed',
            {
                'kind': 'matrix',
                'values': [['Revenue', None]],
                'formulas': [['', '=B2*B3']],
            },
        )


@pytest.mark.parametrize('literal', ['Revenue label', 0, False])
def test_formula_content_does_not_discard_supplied_literal_data(
    tmp_path: Path, literal: str | int | bool
) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('mixed-formulas.xlsx')
    with pytest.raises(ValueError, match='separate'):
        tools.compose_artifact(
            handle,
            'mixed',
            {
                'kind': 'formulas',
                'values': [[literal, None]],
                'formulas': [[None, '=100*2']],
            },
        )


def test_render_requires_placed_content(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('empty.xlsx')
    tools.put_sheet(handle, 'empty', 'Empty', {'anchor': 'end'})
    with pytest.raises(ValueError, match='placed'):
        tools.render_workbook(handle)
    assert not (tmp_path / 'empty.xlsx').exists()


def test_number_format_schema_advertises_supported_choices() -> None:
    schema = TypeAdapter(RangeFormat).json_schema()['properties']['number_format']
    choices = schema.get('anyOf', [schema])
    assert any('$#,##0.00' in choice.get('enum', []) for choice in choices)


def test_strict_provider_can_set_unused_content_fields_to_null(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('nullable.xlsx')
    tools.compose_artifact(
        handle,
        'header',
        {
            'kind': 'row',
            'values': ['Month', 'Revenue'],
            'value': None,
            'formulas': None,
        },
    )
    tools.compose_artifact(
        handle,
        'formula',
        {
            'kind': 'formulas',
            'values': None,
            'value': None,
            'formulas': [['=B2*C2']],
        },
    )


def test_matrix_shape_error_explains_vertical_rows(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('shape.xlsx')
    with pytest.raises(ValueError, match='rows'):
        tools.compose_artifact(handle, 'bad', {'kind': 'matrix', 'values': [100, 120]})


def test_shrinking_composition_releases_cells_for_separate_formats(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('resize.xlsx')
    tools.put_sheet(handle, 'main', 'Main', {'anchor': 'end'})
    labels = tools.compose_artifact(handle, 'labels', {'kind': 'row', 'values': ['Price', 100]})
    tools.put_range(handle, 'labels', 'main', 'A1', labels, {})
    tools.compose_artifact(handle, 'labels', {'kind': 'scalar', 'value': 'Price'})
    value = tools.compose_artifact(handle, 'price', {'kind': 'scalar', 'value': 120})
    tools.put_range(handle, 'price', 'main', 'B1', value, {'number_format': '$#,##0.00'})
    sheet = load_workbook(tools.render_workbook(handle).path)['Main']
    assert sheet['A1'].value == 'Price'
    assert sheet['B1'].value == 120


def test_expanding_composition_refuses_collision_without_mutation(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('collision.xlsx')
    tools.put_sheet(handle, 'main', 'Main', {'anchor': 'end'})
    label = tools.compose_artifact(handle, 'label', {'kind': 'scalar', 'value': 'Price'})
    value = tools.compose_artifact(handle, 'value', {'kind': 'scalar', 'value': 100})
    tools.put_range(handle, 'label', 'main', 'A1', label, {})
    tools.put_range(handle, 'value', 'main', 'B1', value, {})
    before = tools.read_workbook(handle)
    with pytest.raises(ValueError, match='overlap'):
        tools.compose_artifact(handle, 'label', {'kind': 'row', 'values': ['Price', 999]})
    assert tools.read_workbook(handle) == before
