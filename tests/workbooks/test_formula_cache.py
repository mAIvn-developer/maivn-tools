"""Delivered workbooks carry usable results without requiring client recalculation."""

from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from maivn_tools import WorkbooksToolSet
from maivn_tools.connectors.workbooks.calculation import (
    CalculationUnavailableError,
    CellError,
    WorkbookCalculator,
)


def test_delivered_formula_values_and_chart_are_available_before_excel(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    book = tools.create_workbook('forecast.xlsx')
    tools.put_sheet(book, 'forecast', 'Forecast', {'anchor': 'end'})
    source = tools.compose_artifact(
        book,
        'inputs',
        {
            'kind': 'matrix',
            'values': [
                ['Month', 'Users', 'ARPU'],
                ['Sep', 10000, 9.5],
                ['Oct', 10800, 9.7],
                ['Nov', 11600, 9.9],
            ],
        },
    )
    tools.put_range(book, 'inputs', 'forecast', 'A1', source, {})
    header = tools.compose_artifact(book, 'header', {'kind': 'scalar', 'value': 'Revenue'})
    tools.put_range(
        book,
        'header',
        'forecast',
        'D1',
        header,
        {
            'chart': {
                'type': 'line',
                'title': 'Revenue',
                'anchor': 'F2',
                'source_range': 'A1:D4',
                'series_columns': [4],
            }
        },
    )
    formulas = tools.compose_artifact(
        book,
        'formulas',
        {'kind': 'formulas', 'formulas': [['=B2*C2'], ['=B3*C3'], ['=B4*C4'], ['=SUM(D2:D4)']]},
    )
    tools.put_range(book, 'formulas', 'forecast', 'D2', formulas, {'number_format': '$#,##0.00'})
    result = tools.render_workbook(book)
    cached = load_workbook(result.path, data_only=True)['Forecast']
    assert [cached[f'D{row}'].value for row in range(2, 6)] == pytest.approx(
        [95000, 104760, 114840, 314600]
    )
    editable = load_workbook(result.path)['Forecast']
    assert editable['D5'].value == '=SUM(D2:D4)'
    assert editable['D2'].number_format == '$#,##0.00'
    with ZipFile(result.path) as archive:
        chart = ET.fromstring(archive.read('xl/charts/chart1.xml'))  # noqa: S314 - local generated XML.
    namespaces = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
    assert [
        float(node.text or '')
        for node in chart.findall('.//c:val/c:numRef/c:numCache/c:pt/c:v', namespaces)
    ] == pytest.approx([95000, 104760, 114840])


@pytest.mark.parametrize(
    ('formula', 'expected'),
    [
        ('=A1+A2*2', 16),
        ('=(A1+A2)*2', 20),
        ('=-2^2', 4),
        ('=2^3^2', 64),
        ('=2^-2', 0.25),
        ('=A1*(1+15%)', 4.6),
        ('=SUM(A1:A5)', 10),
        ('=AVERAGE(A1:A5)', 5),
        ('=COUNT(A1:A5)', 2),
        ('=MAX(A1:A5)', 6),
        ('=MIN(A1:A5)', 4),
        ('=SUM(A1:A5,MAX(A1:A2),2)', 18),
        ("='Growth assumptions'!$B$2*A1", 8),
        ('=forecast!a1+A2', 10),
        ('=A3', 'label'),
        ('=A4', True),
        ('=A5', 0),
        ('=A4+1', 2),
        ('=A6+1', 13),
        ('=COUNT(A1:A7)', 2),
        ('=A3+1', CellError('#VALUE!')),
        ('=SUM(A1:A7)', CellError('#DIV/0!')),
        ('=1/0', CellError('#DIV/0!')),
        ('=Missing!A1', CellError('#REF!')),
        ('=(-1)^0.5', CellError('#NUM!')),
        ('=AVERAGE(A8:A9)', CellError('#DIV/0!')),
        ('=SUM(A8:A9)', 0),
    ],
)
def test_supported_formula_results_match_spreadsheet_semantics(
    formula: str, expected: object
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = 'Forecast'
    for index, value in enumerate([4, 6, 'label', True, None, '12', '=1/0'], 1):
        sheet.cell(index, 1, value)
    workbook.create_sheet('Growth assumptions')['B2'] = 2
    sheet['B1'] = formula
    result = WorkbookCalculator(workbook).cell('Forecast', 'B1')
    assert result == (pytest.approx(expected) if isinstance(expected, (float, int)) else expected)


def test_dependency_cache_handles_forward_references_and_refuses_cycles() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet['A1'], sheet['A2'], sheet['A3'] = '=A2+1', '=A3*2', 5
    calculator = WorkbookCalculator(workbook)
    assert calculator.cell(sheet.title, 'A1') == 11
    sheet['B1'], sheet['B2'] = '=B2', '=B1'
    with pytest.raises(CalculationUnavailableError):
        calculator.cell(sheet.title, 'B1')
    assert calculator.cell(sheet.title, 'A1') == 11


def test_huge_ranges_do_not_allocate_or_invent_cached_results() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet['A1'] = '=SUM(B1:XFD1048576)'
    with pytest.raises(CalculationUnavailableError):
        WorkbookCalculator(workbook).cell(sheet.title, 'A1')
    assert sheet.max_row == 1


def test_cached_errors_and_strings_round_trip_and_edit_recalculates(tmp_path: Path) -> None:
    tools = WorkbooksToolSet(tmp_path)
    book = tools.create_workbook('types.xlsx')
    tools.put_sheet(book, 'data', 'Data', {'anchor': 'end'})
    inputs = tools.compose_artifact(book, 'inputs', {'kind': 'row', 'values': [5, 'label', True]})
    tools.put_range(book, 'inputs', 'data', 'A1', inputs, {})
    formulas = tools.compose_artifact(
        book,
        'calculations',
        {'kind': 'formulas', 'formulas': [['=A1*2', '=B1', '=C1', '=1/0', '=Missing!A1']]},
    )
    tools.put_range(book, 'formulas', 'data', 'A2', formulas, {})
    original = tools.render_workbook(book)
    cached = load_workbook(original.path, data_only=True)['Data']
    assert [cached.cell(2, column).value for column in range(1, 6)] == [
        10,
        'label',
        True,
        '#DIV/0!',
        '#REF!',
    ]
    assert cached['D2'].data_type == 'e'
    tools.compose_artifact(book, 'inputs', {'kind': 'row', 'values': [12, 'label', True]})
    revised = tools.render_workbook(book, overwrite=True)
    assert load_workbook(revised.path, data_only=True)['Data']['A2'].value == 24
    assert load_workbook(original.path, data_only=True)['Data']['A2'].value == 10
    assert original.sha256 != revised.sha256


def test_long_prefix_expression_defers_without_recursion_failure() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet['A1'] = '=' + '-' * 600 + '1'
    with pytest.raises(CalculationUnavailableError):
        WorkbookCalculator(workbook).cell(sheet.title, 'A1')
