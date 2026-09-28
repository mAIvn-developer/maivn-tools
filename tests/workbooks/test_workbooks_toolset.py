# pyright: strict
from __future__ import annotations

import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any, cast

import pytest
from maivn import Agent
from maivn._internal.compat.decorators import (
    ARGUMENT_PRODUCER_REQUIREMENTS_ATTR,
)
from openpyxl import load_workbook

import maivn_tools
from maivn_tools.connectors.workbooks.models import CellValue, WorkbookHandle


def test_workbook_tools_publish_state_and_composition_dependency_contracts(
    tmp_path: Path,
) -> None:
    tools = Agent(name='workbooks-contract', api_key='test-key').add_toolset(
        maivn_tools.WorkbooksToolSet(tmp_path)
    )
    by_name = {tool.name: tool for tool in tools}

    read_schema = by_name['WORKBOOKS_read_workbook'].output_schema
    assert isinstance(read_schema, dict)
    assert read_schema['required'] == ['workbook', 'compositions', 'sheets', 'readback']
    readback = cast('dict[str, Any]', read_schema['properties']['readback'])
    assert 'actual placed ranges' in readback['description']
    assert not getattr(
        by_name['WORKBOOKS_put_range'].target,
        ARGUMENT_PRODUCER_REQUIREMENTS_ATTR,
        (),
    )


def _build_workbook(toolset: maivn_tools.WorkbooksToolSet, filename: str) -> WorkbookHandle:
    book = toolset.create_workbook(filename)
    toolset.put_sheet(book, 'summary', 'Summary', {'anchor': 'end'}, freeze_panes='A2')
    data = toolset.compose_artifact(
        book,
        'data',
        {'kind': 'matrix', 'values': [['Metric', 'Value'], ['Ready', 1]]},
    )
    formulas = toolset.compose_artifact(
        book, 'formula', {'kind': 'formulas', 'formulas': [['=SUM(B2:B2)']]}
    )
    toolset.put_range(
        book,
        'data-range',
        'summary',
        'A1',
        data,
        {
            'style': 'header_row',
            'number_format': 'General',
            'chart': {'type': 'bar', 'title': 'Readiness', 'anchor': 'D2'},
        },
    )
    toolset.put_range(
        book,
        'total-formula',
        'summary',
        'B3',
        formulas,
        {'style': 'total', 'number_format': '#,##0'},
    )
    return book


def test_workbook_stepwise_chain_reopens_and_is_byte_deterministic(tmp_path: Path) -> None:
    left = maivn_tools.WorkbooksToolSet(tmp_path / 'left')
    right = maivn_tools.WorkbooksToolSet(tmp_path / 'right')
    left_book = _build_workbook(left, 'proof.xlsx')
    right_book = _build_workbook(right, 'proof.xlsx')

    data_ref = left.compose_artifact(
        left_book,
        'data',
        {'kind': 'matrix', 'values': [['Metric', 'Value'], ['Ready', 2]]},
    )
    left.put_range(
        left_book,
        'data-range',
        'summary',
        'A1',
        data_ref,
        {
            'style': 'header_row',
            'number_format': 'General',
            'chart': {'type': 'bar', 'title': 'Readiness', 'anchor': 'D2'},
        },
    )
    left.compose_artifact(
        left_book,
        'data',
        {'kind': 'matrix', 'values': [['Metric', 'Value'], ['Ready', 1]]},
    )

    left_file = left.render_workbook(left_book)
    right_file = right.render_workbook(right_book)
    assert left_file.sha256 == right_file.sha256
    workbook = load_workbook(left_file.path, data_only=False)
    assert workbook.sheetnames == ['Summary']
    assert workbook['Summary']['B2'].value == 1
    assert workbook['Summary']['B3'].value == '=SUM(B2:B2)'
    assert workbook['Summary'].freeze_panes == 'A2'
    assert workbook.calculation.calcMode == 'auto'
    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True
    with zipfile.ZipFile(left_file.path) as archive:
        assert archive.namelist() == sorted(archive.namelist())
        assert 'xl/charts/chart1.xml' in archive.namelist()
        assert {entry.date_time for entry in archive.infolist()} == {(1980, 1, 1, 0, 0, 0)}
        core_properties = archive.read('docProps/core.xml')
        assert b'<dcterms:modified' in core_properties
        assert b'>2000-01-01T00:00:00Z</dcterms:modified>' in core_properties


def test_fresh_instance_edit_preserves_formula_chart_and_formatting(tmp_path: Path) -> None:
    tools = maivn_tools.WorkbooksToolSet(tmp_path)
    handle = _build_workbook(tools, 'proof.xlsx')
    tools.render_workbook(handle)
    reopened = maivn_tools.WorkbooksToolSet(tmp_path)
    reopened.compose_artifact(
        handle,
        'data',
        {
            'kind': 'matrix',
            'values': [['Metric', 'Value'], ['Ready', 99]],
        },
    )
    result = reopened.render_workbook(handle, overwrite=True)
    book = load_workbook(result.path, data_only=False)
    assert book['Summary']['B2'].value == 99
    assert book['Summary']['B3'].value == '=SUM(B2:B2)'
    assert book['Summary']['B3'].number_format == '#,##0'
    assert book['Summary'].freeze_panes == 'A2'
    with zipfile.ZipFile(result.path) as archive:
        assert 'xl/charts/chart1.xml' in archive.namelist()


def test_print_fit_roundtrip_and_wide_sheet_readability(tmp_path: Path) -> None:
    tools = maivn_tools.WorkbooksToolSet(tmp_path)
    handle = _build_workbook(tools, 'print.xlsx')
    tools.put_sheet(handle, 'wide', 'Wide', {'anchor': 'end'})
    ref = tools.compose_artifact(
        handle, 'wide', {'kind': 'row', 'values': cast('list[CellValue]', ['Header'] * 60)}
    )
    tools.put_range(handle, 'wide', 'wide', 'A1', ref, {'style': 'header_row'})
    result = tools.render_workbook(handle)
    book = load_workbook(result.path)
    assert book['Summary'].page_setup.fitToWidth == 1
    assert book['Summary'].page_setup.fitToHeight == 0
    assert book['Wide'].page_setup.scale == 100
    page_properties = book['Wide'].sheet_properties.pageSetUpPr
    assert page_properties is not None
    assert not page_properties.fitToPage
    tools.put_sheet(
        handle,
        'summary',
        'Summary',
        {'anchor': 'start'},
        print_settings={'mode': 'fit_width', 'orientation': 'landscape'},
    )
    reopened = maivn_tools.WorkbooksToolSet(tmp_path)
    reopened.put_sheet(handle, 'summary', 'Renamed', {'anchor': 'start'})
    result = reopened.render_workbook(handle, overwrite=True)
    book = load_workbook(result.path)
    assert book['Renamed'].page_setup.orientation == 'landscape'
    assert book['Renamed'].page_setup.fitToWidth == 1
    assert book['Renamed'].page_setup.fitToHeight == 0


def test_explicit_print_fit_refuses_unreadable_scale(tmp_path: Path) -> None:
    tools = maivn_tools.WorkbooksToolSet(tmp_path)
    handle = tools.create_workbook('wide.xlsx')
    tools.put_sheet(handle, 'wide', 'Wide', {'anchor': 'end'}, print_settings={'mode': 'fit_width'})
    ref = tools.compose_artifact(
        handle, 'wide', {'kind': 'row', 'values': cast('list[CellValue]', ['Header'] * 60)}
    )
    tools.put_range(handle, 'wide', 'wide', 'A1', ref, {})
    with pytest.raises(ValueError, match='minimum print scale'):
        tools.render_workbook(handle)
    assert not (tmp_path / 'wide.xlsx').exists()


def test_workbook_refuses_inline_content_overlap_and_unsafe_paths(tmp_path: Path) -> None:
    books = maivn_tools.WorkbooksToolSet(tmp_path)
    book = books.create_workbook('proof.xlsx')
    books.put_sheet(book, 'sheet', 'Sheet', {'anchor': 'end'})
    ref = books.compose_artifact(book, 'first', {'kind': 'matrix', 'values': [[1, 2], [3, 4]]})
    books.put_range(book, 'first', 'sheet', 'A1', ref, {'style': 'body'})
    with pytest.raises(ValueError, match='overlap') as failure:
        books.put_range(book, 'second', 'sheet', 'B2', ref, {'style': 'body'})
    assert getattr(type(failure.value), 'sdk_error_code', None) == 'sdk_workbook_range_overlap'
    # Replacing the same range remains supported; overlapping a different range does not.
    books.put_range(book, 'first', 'sheet', 'A1', ref, {'style': 'body'})
    with pytest.raises(maivn_tools.CompositionDependencyError, match='WORKBOOKS_compose_artifact'):
        books.put_range(book, 'inline', 'sheet', 'D1', 'raw values', {'style': 'body'})  # type: ignore[arg-type]
    with pytest.raises(ValueError, match='directly beneath output_dir'):
        books.create_workbook('../escape.xlsx')


def test_workbook_read_remove_and_one_shot_route_through_compositions(tmp_path: Path) -> None:
    books = maivn_tools.WorkbooksToolSet(tmp_path)
    generated = books.create_xlsx(
        'small.xlsx',
        {
            'sheets': [
                {
                    'sheet_id': 'main',
                    'name': 'Main',
                    'ranges': [
                        {
                            'range_id': 'values',
                            'start_cell': 'A1',
                            'content': {'kind': 'row', 'values': ['A', 'B']},
                        }
                    ],
                }
            ]
        },
    )
    handle = cast('WorkbookHandle', generated.workspace)
    state = cast('dict[str, Any]', books.read_workbook(handle))
    assert state['sheets'][0]['ranges'][0]['range_id'] == 'values'
    assert state['compositions'][0]['composition_id'] == 'main-values-content'
    assert books.remove_range(handle, 'values')['removed'] is True


def test_workbook_handles_are_isolated_under_concurrency(tmp_path: Path) -> None:
    books = maivn_tools.WorkbooksToolSet(tmp_path)

    def create_handle(_: int) -> WorkbookHandle:
        return books.create_workbook('same.xlsx')

    with ThreadPoolExecutor(max_workers=4) as executor:
        handles = list(executor.map(create_handle, range(8)))

    assert len({handle['workbook_id'] for handle in handles}) == 8
