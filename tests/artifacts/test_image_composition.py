from __future__ import annotations

import asyncio
import base64
import hashlib
import io
import json
import zipfile
from pathlib import Path
from typing import Any, cast

import httpx
import pytest
from docx import Document
from maivn import Agent, Client, ClientConfig
from maivn._internal.artifact_image_context import InvocationArtifactImages
from maivn._internal.models import StreamEvent
from maivn_contracts.artifacts import OrdinaryArtifactRef, PrivateArtifactRef
from PIL import Image
from PIL.PngImagePlugin import PngInfo
from pptx import Presentation
from pydantic import AnyUrl
from pypdf import PdfReader

import maivn_tools
from maivn_tools.connectors.documents.models import DocumentBlock
from tests.artifacts.test_revision_sources import (
    _private_receipt,  # pyright: ignore[reportPrivateUsage] - canonical private source fixture.
    _receipt,  # pyright: ignore[reportPrivateUsage] - canonical ordinary source fixture.
)


def image_content(format_name: str = 'PNG') -> dict[str, object]:
    stream = io.BytesIO()
    info = PngInfo()
    info.add_text('source', 'C:/private/source.png')
    Image.new('RGB', (160, 80), '#2479ab').save(stream, format=format_name, pnginfo=info)
    return {
        'kind': 'image',
        'image': {
            'content_base64': base64.b64encode(stream.getvalue()).decode('ascii'),
            'mime_type': 'image/png' if format_name == 'PNG' else 'image/jpeg',
            'alt_text': 'Blue rectangle',
        },
    }


def test_registered_placement_accepts_saved_image_but_rejects_foreign_or_missing_reference(
    tmp_path: Path,
) -> None:
    """A fresh invocation can place an existing image without a text producer call."""
    original = maivn_tools.DocumentsToolSet(tmp_path)
    handle = original.create_document('saved-image.docx')
    image_reference = original.compose_image(
        handle, 'visual', **cast('dict[str, Any]', image_content()['image'])
    )
    restored = maivn_tools.DocumentsToolSet(tmp_path)
    tools = Agent(name='restored-image', api_key='test-key').add_toolset(restored)
    placement = next(tool for tool in tools if tool.name == 'DOCUMENTS_put_block')
    assert placement.target is not None
    placement.target(
        doc=handle,
        block_id='visual',
        block={
            'type': 'image',
            'composition': image_reference,
            'image_placement': {'width_inches': 4, 'height_inches': 2, 'alignment': 'center'},
        },
        position={'anchor': 'end'},
    )
    rendered = restored.render_document(handle)
    assert len(Document(rendered.path).inline_shapes) == 1
    for invalid_reference in (
        {**image_reference, 'document_id': '0' * 32},
        {**image_reference, 'composition_id': 'invented'},
    ):
        with pytest.raises(maivn_tools.CompositionDependencyError):
            placement.target(
                doc=handle,
                block_id='bad',
                block={'type': 'image', 'composition': invalid_reference},
                position={'anchor': 'end'},
            )
    assert len(restored.read_document(handle)['outline']) == 1


@pytest.mark.parametrize('format_name', ['PNG', 'JPEG'])
@pytest.mark.parametrize('kind', ['documents', 'pdf', 'presentations'])
def test_image_survives_fresh_instance_text_edit(
    tmp_path: Path, kind: str, format_name: str
) -> None:
    factory = {
        'documents': maivn_tools.DocumentsToolSet,
        'pdf': maivn_tools.PDFToolSet,
        'presentations': maivn_tools.PresentationsToolSet,
    }[kind]
    tools = cast('Any', factory(tmp_path))
    if kind == 'presentations':
        handle = tools.create_presentation('proof.pptx')
        tools.put_slide(handle, 'slide', {'anchor': 'end'})
    else:
        handle = tools.create_document('proof.docx' if kind == 'documents' else 'proof.pdf')
    image_ref = tools.compose_image(
        handle, 'visual', **cast('dict[str, Any]', image_content(format_name)['image'])
    )
    text_ref = tools.compose_artifact(handle, 'body', {'kind': 'text', 'text': 'Original'})
    if kind == 'presentations':
        tools.put_element(
            handle,
            'slide',
            'visual',
            {
                'type': 'image',
                'composition': image_ref,
                'x': 1,
                'y': 2,
                'width': 4,
                'height': 4,
            },
        )
        tools.put_element(
            handle,
            'slide',
            'body',
            {
                'type': 'text',
                'composition': text_ref,
                'x': 1,
                'y': 0.5,
                'width': 4,
                'height': 1,
                'font_size': 24,
                'alignment': 'center',
            },
        )
        first = tools.render_presentation(handle)
    else:
        tools.put_block(
            handle,
            'visual',
            {
                'type': 'image',
                'composition': image_ref,
                'image_placement': {'width_inches': 4, 'height_inches': 4, 'alignment': 'center'},
            },
            {'anchor': 'end'},
        )
        tools.put_block(
            handle, 'body', {'type': 'paragraph', 'composition': text_ref}, {'anchor': 'end'}
        )
        first = tools.render_document(handle)
    original = Path(first.path).read_bytes()
    reopened = cast('Any', factory(tmp_path))
    reopened.compose_artifact(handle, 'body', {'kind': 'text', 'text': 'Revised text'})
    result = (
        reopened.render_presentation(handle, overwrite=True)
        if kind == 'presentations'
        else reopened.render_document(handle, overwrite=True)
    )
    assert result.kind == 'generated_file'
    assert Path(result.path).read_bytes() != original
    if kind == 'documents':
        doc = Document(result.path)
        assert [p.text for p in doc.paragraphs][-1] == 'Revised text'
        assert len(doc.inline_shapes) == 1
        assert (
            cast('Any', doc.inline_shapes[0]).width / cast('Any', doc.inline_shapes[0]).height == 2
        )
    elif kind == 'pdf':
        pdf = PdfReader(result.path)
        assert 'Revised text' in pdf.pages[0].extract_text()
        assert len(pdf.pages[0].images) == 1
    else:
        deck = Presentation(result.path)
        shapes = cast('Any', deck.slides[0].shapes)
        assert shapes[0].image.size == (160, 80)
        assert shapes[0].width / shapes[0].height == 2
        assert shapes[1].text == 'Revised text'
        assert shapes[1].text_frame.paragraphs[0].font.size.pt == 24
    if kind != 'pdf':
        with zipfile.ZipFile(result.path) as archive:
            media = [name for name in archive.namelist() if '/media/' in name]
            assert len(media) == 1
            assert b'C:/private/source.png' not in archive.read(media[0])
            with zipfile.ZipFile(io.BytesIO(original)) as before:
                assert before.read(media[0]) == archive.read(media[0])


@pytest.mark.parametrize(
    'bad_image',
    [
        {'content_base64': 'not base64', 'mime_type': 'image/png', 'alt_text': 'x'},
        {
            'content_base64': base64.b64encode(b'<svg/>').decode(),
            'mime_type': 'image/png',
            'alt_text': 'x',
        },
    ],
)
def test_invalid_image_refused_without_persisting(
    tmp_path: Path, bad_image: dict[str, str]
) -> None:
    tools = maivn_tools.DocumentsToolSet(tmp_path)
    handle = tools.create_document('proof.docx')
    with pytest.raises(ValueError, match=r'image|base64'):
        tools.compose_image(handle, 'bad', **cast('Any', bad_image))
    assert tools.read_document(handle)['compositions'] == []


def test_slide_layout_reports_overflow_and_rejects_nonfinite_geometry(tmp_path: Path) -> None:
    tools = cast('Any', maivn_tools.PresentationsToolSet(tmp_path))
    handle = tools.create_presentation('proof.pptx')
    tools.put_slide(handle, 'slide', {'anchor': 'end'})
    reference = tools.compose_artifact(handle, 'text', {'kind': 'text', 'text': 'Long text ' * 200})
    element = {
        'type': 'text',
        'composition': reference,
        'x': 1,
        'y': 1,
        'width': 2,
        'height': 0.4,
        'font_size': 28,
    }
    tools.put_element(handle, 'slide', 'body', element)
    issues = tools.inspect_layout(handle)
    assert any(
        issue['code'] == 'text_overflow' and issue['element_id'] == 'body' for issue in issues
    )
    with pytest.raises(ValueError, match='finite'):
        tools.put_element(handle, 'slide', 'invalid', {**element, 'x': float('nan')})


@pytest.mark.parametrize(
    'factory',
    [maivn_tools.DocumentsToolSet, maivn_tools.PDFToolSet, maivn_tools.PresentationsToolSet],
)
def test_image_admission_is_local_and_readback_contains_only_reference(
    tmp_path: Path,
    factory: Any,
) -> None:
    toolset = factory(tmp_path)
    if factory == maivn_tools.PresentationsToolSet:
        handle = toolset.create_presentation('proof.pptx')
        read = toolset.read_presentation
    else:
        handle = toolset.create_document(
            'proof.docx' if factory == maivn_tools.DocumentsToolSet else 'proof.pdf'
        )
        read = toolset.read_document
    toolset.compose_image(handle, 'visual', **cast('dict[str, Any]', image_content()['image']))
    state = json.dumps(read(handle))
    assert 'asset_id' in state
    assert 'content_base64' not in state
    assert 'iVBOR' not in state
    published = Agent(name='images', api_key='test-key').add_toolset(toolset)
    assert all(not tool.name.endswith('compose_image') for tool in published)
    assert 'content_base64' not in json.dumps([tool.input_schema for tool in published])


def test_image_mime_size_metadata_and_missing_placement_refused(tmp_path: Path) -> None:
    tools = maivn_tools.DocumentsToolSet(tmp_path)
    handle = tools.create_document('proof.docx')
    image = cast('dict[str, Any]', image_content()['image'])
    with pytest.raises(ValueError, match='MIME'):
        tools.compose_image(handle, 'wrong', **{**image, 'mime_type': 'image/jpeg'})
    with pytest.raises(ValueError, match='alt_text'):
        tools.compose_image(handle, 'wrong', **{**image, 'alt_text': 'x' * 1001})
    with pytest.raises(ValueError, match='byte limit'):
        tools.compose_image(handle, 'wrong', **{**image, 'content_base64': 'a' * 12_000_000})
    reference = tools.compose_image(handle, 'visual', **image)
    with pytest.raises(ValueError, match='image_placement'):
        tools.put_block(
            handle,
            'bad',
            cast('DocumentBlock', {'type': 'image', 'composition': reference}),
            {'anchor': 'end'},
        )
    with pytest.raises(ValueError, match='fit within'):
        tools.put_block(
            handle,
            'bad',
            {
                'type': 'image',
                'composition': reference,
                'image_placement': {'width_inches': 100, 'height_inches': 1, 'alignment': 'left'},
            },
            {'anchor': 'end'},
        )


def test_unknown_manifest_version_refused(tmp_path: Path) -> None:
    tools = maivn_tools.DocumentsToolSet(tmp_path)
    handle = tools.create_document('proof.docx')
    source = tmp_path / '.maivn-documents' / f'{handle["document_id"]}.json'
    manifest = json.loads(source.read_text())
    manifest['version'] = 999
    source.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match='manifest version'):
        maivn_tools.DocumentsToolSet(tmp_path).read_document(handle)


def test_reused_image_reference_validates_metadata(tmp_path: Path) -> None:
    tools = maivn_tools.DocumentsToolSet(tmp_path)
    handle = tools.create_document('proof.docx')
    tools.compose_image(handle, 'visual', **cast('dict[str, Any]', image_content()['image']))
    reference = tools.read_document(handle)['compositions'][0]['content'].get('image')
    assert reference is not None
    with pytest.raises(ValueError, match='alt_text'):
        tools.compose_artifact(
            handle,
            'bad',
            {
                'kind': 'image',
                'image': {**reference, 'alt_text': '\x00'},
            },
        )


@pytest.mark.parametrize('variant', ['too_wide', 'animated'])
def test_image_dimension_and_animation_limits(tmp_path: Path, variant: str) -> None:
    output = io.BytesIO()
    if variant == 'too_wide':
        Image.new('RGB', (8193, 1)).save(output, format='PNG')
    else:
        Image.new('RGB', (10, 10), 'red').save(
            output,
            format='PNG',
            save_all=True,
            append_images=[Image.new('RGB', (10, 10), 'blue')],
        )
    tools = maivn_tools.DocumentsToolSet(tmp_path)
    handle = tools.create_document('proof.docx')
    with pytest.raises(ValueError, match=r'limits|frame'):
        tools.compose_image(
            handle,
            'bad',
            content_base64=base64.b64encode(output.getvalue()).decode(),
            mime_type='image/png',
            alt_text='Example',
        )


def test_image_tampering_refused_before_output_write(tmp_path: Path) -> None:
    tools = maivn_tools.DocumentsToolSet(tmp_path)
    handle = tools.create_document('proof.docx')
    reference = tools.compose_image(
        handle, 'visual', **cast('dict[str, Any]', image_content()['image'])
    )
    tools.put_block(
        handle,
        'visual',
        {
            'type': 'image',
            'composition': reference,
            'image_placement': {'width_inches': 2, 'height_inches': 2, 'alignment': 'left'},
        },
        {'anchor': 'end'},
    )
    source = tmp_path / '.maivn-documents' / f'{handle["document_id"]}.json'
    manifest = json.loads(source.read_text())
    asset_id = next(iter(manifest['images']))
    manifest['images'][asset_id] = base64.b64encode(b'changed').decode()
    source.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match='digest mismatch'):
        maivn_tools.DocumentsToolSet(tmp_path).render_document(handle)
    assert not (tmp_path / 'proof.docx').exists()


@pytest.mark.parametrize('private', [False, True])
@pytest.mark.parametrize('kind', ['documents', 'pdf', 'presentations'])
def test_generated_image_reference_composes_and_survives_source_roundtrip(
    tmp_path: Path,
    kind: str,
    *,
    private: bool,
) -> None:
    raw = base64.b64decode(cast('dict[str, str]', image_content()['image'])['content_base64'])
    common = {
        'artifact_id': 'generated-image',
        'logical_output_id': 'image-output',
        'revision': 1,
        'kind': 'image',
        'mime_type': 'image/png',
        'created_at': '2026-09-05T12:00:00Z',
        'effective_retention': {
            'policy_snapshot_id': 'policy',
            'retention_class': 'artifact_30d',
            'expires_at': '2026-10-05T12:00:00Z',
        },
    }
    if private:
        selected = PrivateArtifactRef.model_validate(
            {
                **common,
                'custody': 'vault_private',
                'state': 'available_in_vault',
                'display_label': 'Private image',
                'creation_receipt_id': 'receipt',
                'producer': {
                    'producer_class': 'vault',
                    'producer_id': 'vault',
                    'root_invocation_id': 'root',
                },
                'retrieval_action': {'relation': 'artifact.vault_download_authorization'},
            }
        )
    else:
        selected = OrdinaryArtifactRef.model_validate(
            {
                **common,
                'custody': 'ordinary',
                'state': 'available',
                'display_filename': 'generated.png',
                'size_bytes': len(raw),
                'sha256': hashlib.sha256(raw).hexdigest(),
                'producer': {
                    'producer_class': 'hosted_tool',
                    'producer_id': 'image-generation',
                    'root_invocation_id': 'root',
                },
                'validation_receipt': {
                    'receipt_id': 'receipt',
                    'status': 'validated',
                    'validator_profile': 'image',
                    'validator_version': '1',
                },
                'safe_preview': {'kind': 'image', 'width_pixels': 160, 'height_pixels': 80},
                'retrieval_action': {'relation': 'artifact.download_authorization'},
            }
        )
    downloads: list[str] = []

    def transport(request: httpx.Request) -> httpx.Response:
        assert not private
        assert request.headers['authorization'] == 'Bearer test-key'
        assert request.url.path == '/v1/artifacts/generated-image/download'
        assert dict(request.url.params) == {'revision': '1'}
        downloads.append('ordinary')
        return httpx.Response(
            200,
            content=raw,
            headers={
                'content-type': 'image/png',
                'content-length': str(len(raw)),
                'content-disposition': 'attachment; filename="generated.png"',
            },
        )

    class Private:
        async def adownload(self, ref: object, *, session_id: str) -> bytes:
            assert private
            assert ref == selected
            assert session_id == 'current-session'
            downloads.append('private')
            return raw

    client = Client(
        config=ClientConfig(api_key='test-key', base_url=AnyUrl('http://testserver')),
        transport=httpx.MockTransport(transport),
    )
    registry = InvocationArtifactImages(
        ordinary=client.artifacts, private=cast('Any', Private()), session_id='current-session'
    )
    registry.observe(
        StreamEvent(
            position=1,
            event_type='system_tool_complete',
            data={
                'type': 'system_tool_complete',
                'payload': {
                    'outcome': {
                        'call_id': 'image-generation',
                        'status': 'ok',
                        'duration_ms': 1,
                        'result': {},
                        'artifact_refs': [selected.model_dump(mode='json')],
                    }
                },
            },
        )
    )
    factory, extension = {
        'documents': (maivn_tools.DocumentsToolSet, 'docx'),
        'pdf': (maivn_tools.PDFToolSet, 'pdf'),
        'presentations': (maivn_tools.PresentationsToolSet, 'pptx'),
    }[kind]
    tools = cast('Any', factory(tmp_path / 'first-host'))
    handle = (
        tools.create_presentation('proof.pptx')
        if kind == 'presentations'
        else tools.create_document(f'proof.{extension}')
    )

    async def compose() -> object:
        with registry.activate():
            return await tools.compose_artifact_image(
                handle, 'visual', 'generated-image', 'Blue rectangle'
            )

    composed = asyncio.run(compose())
    assert downloads == ['private' if private else 'ordinary']
    assert 'content_base64' not in json.dumps(composed)
    assert hashlib.sha256(raw).hexdigest() not in json.dumps(composed)
    if kind == 'presentations':
        tools.put_slide(handle, 'slide', {'anchor': 'end'})
        tools.put_element(
            handle,
            'slide',
            'visual',
            {'type': 'image', 'composition': composed, 'x': 1, 'y': 1, 'width': 4, 'height': 2},
            {'anchor': 'end'},
        )
        generated = tools.render_presentation(handle)
    else:
        tools.put_block(
            handle,
            'visual',
            {
                'type': 'image',
                'composition': composed,
                'image_placement': {'width_inches': 4, 'height_inches': 2, 'alignment': 'center'},
            },
            {'anchor': 'end'},
        )
        generated = tools.render_document(handle)
    assert generated.custody == ('vault_private' if private else 'ordinary')
    source = tools.export_generated_file_source(generated)
    source_manifest = json.loads(source)['manifest']
    if private:
        assert next(iter(source_manifest['images'])).startswith('private-image-')
        assert source_manifest['source_custody'] == 'vault_private'
    # Fresh-host restoration uses the generated parent's receipt, not the image receipt.

    parent_ref = _private_receipt(generated) if private else _receipt(generated)
    fresh = cast('Any', factory(tmp_path / 'fresh-host'))
    resumed = (
        fresh.resume_private_artifact(parent_ref, source=source)
        if private
        else fresh.resume_artifact(parent_ref, source=source)
    )
    restored = (
        fresh.render_presentation(resumed)
        if kind == 'presentations'
        else fresh.render_document(resumed)
    )
    assert Path(restored.path).read_bytes() == Path(generated.path).read_bytes()
    assert restored.custody == generated.custody
