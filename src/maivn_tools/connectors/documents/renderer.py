"""Deterministic DOCX byte rendering from persisted document state."""

# pyright: strict

from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone
from typing import Any, cast

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement  # pyright: ignore[reportUnknownVariableType]
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.styles.style import ParagraphStyle

from ..artifacts.images import ImagePlacement, ImageReference, contain_size, image_bytes
from .models import CompositionContent, DocumentBlock
from .layout import table_column_widths

_CANONICAL_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_CANONICAL_DOCUMENT_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


def _apply_document_defaults(document: DocumentObject) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(0.8)
    sizes = {'Normal': 11, 'Title': 26, 'List Bullet': 11, 'List Number': 11}
    sizes.update({f'Heading {level}': max(11, 18 - level * 2) for level in range(1, 10)})
    for name, size in sizes.items():
        style = cast('ParagraphStyle', document.styles[name])
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('000000')
        style.font.bold = name == 'Title' or name.startswith('Heading')
        # Theme references and title borders in the stock template differ across viewers.
        for fonts in style.element.xpath('./w:rPr/w:rFonts'):
            for key in list(fonts.attrib):
                if key.endswith('Theme'):
                    del fonts.attrib[key]
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        paragraph = style.paragraph_format
        paragraph.line_spacing = 1.15
        paragraph.space_after = Pt(8)
        paragraph.widow_control = True
        if name == 'Title' or name.startswith('Heading'):
            paragraph.keep_with_next = True
            paragraph.space_before = Pt(0 if name == 'Title' else 14)
            paragraph.space_after = Pt(12 if name == 'Title' else 6)


def _add_table(document: DocumentObject, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.autofit = False
    widths = table_column_widths(rows, 6.5)
    for column, width in zip(table.columns, widths, strict=True):
        column.width = Inches(width)
    borders = cast('Any', OxmlElement('w:tblBorders'))
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = cast('Any', OxmlElement(f'w:{edge}'))
        for key, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
            border.set(qn(f'w:{key}'), value)
        borders.append(border)
    cast('Any', table)._tbl.tblPr.append(borders)  # noqa: SLF001
    for index, values in enumerate(rows):
        row = table.rows[index]
        properties = cast('Any', row)._tr.get_or_add_trPr()  # noqa: SLF001
        properties.append(OxmlElement('w:cantSplit'))
        if index == 0:
            properties.append(OxmlElement('w:tblHeader'))
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            fill = cast('Any', OxmlElement('w:shd'))
            fill.set(
                qn('w:fill'), 'DCE8F2' if index == 0 else ('F5F7FA' if index % 2 == 0 else 'FFFFFF')
            )
            cast('Any', cell)._tc.get_or_add_tcPr().append(fill)  # noqa: SLF001
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = index == 0
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def render_docx_bytes(manifest: dict[str, object]) -> bytes:
    """Render one manifest into byte-stable DOCX content."""
    document = Document()
    _apply_document_defaults(document)
    properties = document.core_properties
    properties.title = ''
    properties.subject = ''
    properties.author = ''
    properties.keywords = ''
    properties.comments = ''
    properties.category = ''
    properties.identifier = ''
    properties.language = ''
    properties.version = ''
    properties.last_modified_by = ''
    properties.created = _CANONICAL_DOCUMENT_TIME
    properties.modified = _CANONICAL_DOCUMENT_TIME
    properties.last_printed = _CANONICAL_DOCUMENT_TIME
    properties.revision = 1

    compositions = cast('dict[str, CompositionContent]', manifest['compositions'])
    outline = cast('list[dict[str, object]]', manifest['blocks'])
    for entry in outline:
        block = cast('DocumentBlock', entry['block'])
        block_type = block['type']
        if block_type == 'page_break':
            document.add_page_break()
            continue
        reference = block.get('composition')
        if reference is None:
            raise ValueError(f"block type '{block_type}' is missing its composition reference")
        content = compositions[reference['composition_id']]
        if block_type == 'image':
            image = cast('ImageReference', content.get('image'))
            raw = image_bytes(manifest, image)
            placement = cast('ImagePlacement', block.get('image_placement'))
            width, height = contain_size(raw, placement['width_inches'], placement['height_inches'])
            paragraph = document.add_paragraph()
            paragraph.alignment = {
                'left': WD_ALIGN_PARAGRAPH.LEFT,
                'center': WD_ALIGN_PARAGRAPH.CENTER,
                'right': WD_ALIGN_PARAGRAPH.RIGHT,
            }[placement['alignment']]
            picture = paragraph.add_run().add_picture(
                io.BytesIO(raw), Inches(width), Inches(height)
            )
            # python-docx exposes no public alternative-text setter.
            cast('Any', picture)._inline.docPr.set('descr', image['alt_text'])  # noqa: SLF001
        elif block_type == 'title':
            document.add_heading(cast('str', content.get('text')), level=0)
        elif block_type == 'heading':
            document.add_heading(
                cast('str', content.get('text')),
                level=cast('int', block.get('level')),
            )
        elif block_type == 'paragraph':
            document.add_paragraph(cast('str', content.get('text')))
        elif block_type in {'bullet_list', 'numbered_list'}:
            style = 'List Bullet' if block_type == 'bullet_list' else 'List Number'
            for item in cast('list[str]', content.get('items')):
                document.add_paragraph(item, style=style)
        elif block_type == 'table':
            rows = cast('list[list[str]]', content.get('rows'))
            _add_table(document, rows)

    uncooked = io.BytesIO()
    document.save(uncooked)
    return _canonicalize_docx(uncooked.getvalue())


def _canonicalize_docx(raw: bytes) -> bytes:
    source = io.BytesIO(raw)
    destination = io.BytesIO()
    with (
        zipfile.ZipFile(source, 'r') as archive,
        zipfile.ZipFile(
            destination,
            'w',
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as canonical,
    ):
        for name in sorted(archive.namelist()):
            info = zipfile.ZipInfo(name, date_time=_CANONICAL_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            info.flag_bits = 0
            canonical.writestr(info, archive.read(name), compresslevel=9)
    return destination.getvalue()
