"""Step-wise deterministic XLSX generation toolset."""

# pyright: strict

from __future__ import annotations

import base64
import hashlib
import os
import re
from collections.abc import Mapping
from typing import TYPE_CHECKING, LiteralString, cast, get_args

from maivn import tool_output, toolify, toolset
from maivn_contracts.artifacts import GeneratedFile, OrdinaryArtifactRef, PrivateArtifactRef
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from pydantic import TypeAdapter, ValidationError
from pydantic_core import PydanticCustomError

from ...core.metadata import AuthMode, ProviderCapability, ProviderMetadata
from ...core.permissions import PermissionFlag, PermissionSet
from ..artifacts import ManifestWorkspace
from ..artifacts.revisions import (
    bind_private_receipt,
    expected_private_base,
    resume_private_source,
    immutable_render_path,
    bind_receipt,
    expected_base,
    resume_source,
    snapshot_source,
    source_custody,
    export_source,
    import_source,
)
from ..documents.dependencies import CompositionDependencyError, requires_composition
from .formulas import validate_formula
from .models import (
    ExcelCell,
    ExcelNumberFormat,
    OneShotWorkbook,
    RangeFormat,
    SheetPrintSettings,
    WorkbookCompositionContent,
    WorkbookCompositionReference,
    WorkbookCompositionState,
    WorkbookFormulaAddress,
    WorkbookHandle,
    WorkbookPlacedRangeReadback,
    WorkbookPosition,
    WorkbookReadback,
    WorkbookSheetReadback,
    WorkbookState,
    WorkbookUnplacedCompositionReadback,
)
from .renderer import chart_series_columns, content_matrix, render_xlsx_bytes

if TYPE_CHECKING:
    from pathlib import Path

_CONTENT = TypeAdapter(WorkbookCompositionContent)
_POSITION = TypeAdapter(WorkbookPosition)
_FORMAT = TypeAdapter(RangeFormat)
_PRINT_SETTINGS = TypeAdapter(SheetPrintSettings)
_ONE_SHOT = TypeAdapter(OneShotWorkbook)
_HANDLE_SCHEMA = TypeAdapter(WorkbookHandle).json_schema()
_REFERENCE_SCHEMA = TypeAdapter(WorkbookCompositionReference).json_schema()
_STATE_SCHEMA = TypeAdapter(WorkbookState).json_schema()
_GENERATED_SCHEMA = GeneratedFile.model_json_schema()
_XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
_NUMBER_FORMATS = frozenset(get_args(ExcelNumberFormat))
_READBACK_MAX_SHEETS = 50
_READBACK_MAX_RANGES = 100
_READBACK_MAX_FORMULA_ADDRESSES = 100
_READBACK_MAX_UNPLACED_COMPOSITIONS = 100


def _content_error(
    rule: LiteralString, message: LiteralString, *location: str | int
) -> ValidationError:
    """Mint static rule/location feedback without retaining workbook cell values."""
    return ValidationError.from_exception_data(
        'WorkbookCompositionContent',
        [
            {
                'type': PydanticCustomError(rule, message),
                'loc': ('content', *location),
                'input': None,
            }
        ],
        hide_input=True,
    )


class WorkbookRangeOverlapError(ValueError):
    """A range update must replace its existing placement or use unoccupied cells."""

    sdk_error_code = 'sdk_workbook_range_overlap'


@toolset(prefix='workbooks', metadata={'serialize_calls': True})
class WorkbooksToolSet:
    metadata = ProviderMetadata(
        name='workbooks',
        display_name='Workbooks',
        version='0.1.0',
        description='Compose, inspect, and deterministically render local XLSX files.',
        auth_modes=(AuthMode.NONE,),
        capabilities=frozenset({ProviderCapability.READ, ProviderCapability.WRITE}),
        tags=('artifacts', 'spreadsheets', 'local'),
    )

    def __init__(self, output_dir: str | os.PathLike[str]) -> None:
        self._workspace = ManifestWorkspace(
            output_dir,
            workspace_name='workbooks',
            id_field='workbook_id',
            extension='.xlsx',
        )
        self.connection = None

    def export_generated_file_source(self, generated: GeneratedFile) -> bytes:
        """Export the exact portable source through SDK-controlled artifact custody."""
        return export_source(self._workspace.workspace_dir, 'workbooks', generated)

    def bind_generated_file_receipt(
        self, generated: GeneratedFile, artifact: OrdinaryArtifactRef
    ) -> None:
        """Persist an SDK-verified receipt and its exact rendered source."""
        bind_receipt(self._workspace.workspace_dir, 'workbook_id', generated, artifact)

    def resume_artifact(
        self, artifact: OrdinaryArtifactRef, *, source: bytes | None = None
    ) -> WorkbookHandle:
        """Resume the exact published source retained in this output directory.

        The selected artifact is the expected base for the next render. The
        data plane independently authorizes the caller and conversation.
        """
        if source is not None:
            import_source(self._workspace.workspace_dir, 'workbooks', artifact, source)
        return cast(
            'WorkbookHandle', resume_source(self._workspace.workspace_dir, 'workbook_id', artifact)
        )

    def bind_private_generated_file_receipt(
        self, generated: GeneratedFile, artifact: PrivateArtifactRef
    ) -> None:
        """Persist the SDK-verified private revision for subsequent renders."""
        bind_private_receipt(self._workspace.workspace_dir, 'workbook_id', generated, artifact)

    def resume_private_artifact(
        self, artifact: PrivateArtifactRef, *, source: bytes
    ) -> WorkbookHandle:
        """Resume the exact bytes from private_artifacts.download_source on any host."""
        return cast(
            'WorkbookHandle',
            resume_private_source(
                self._workspace.workspace_dir, 'workbooks', 'workbook_id', artifact, source
            ),
        )

    def sdk_private_workspace(self, arguments: Mapping[str, object]) -> bool:
        """Let the SDK shield readback from this exact locally owned private source."""
        handle = arguments.get('workbook')
        if not isinstance(handle, Mapping):
            return False
        with self._workspace.lock:
            return (
                source_custody(self._workspace.load(cast('Mapping[str, object]', handle)))
                == 'vault_private'
            )

    def authorized_generated_file_roots(self) -> tuple[Path, ...]:
        """Return the output roots the SDK may admit for this toolset."""
        return (self._workspace.output_dir,)

    @toolify(permissions=PermissionSet(PermissionFlag.WRITE))
    @tool_output(_HANDLE_SCHEMA)
    def create_workbook(self, filename: str) -> WorkbookHandle:
        """Start an empty workbook for stepwise construction; this does not publish a file."""
        return cast(
            'WorkbookHandle',
            self._workspace.create(filename, {'compositions': {}, 'sheets': []}),
        )

    @toolify(permissions=PermissionSet(PermissionFlag.WRITE))
    @tool_output(_REFERENCE_SCHEMA)
    def compose_artifact(
        self,
        workbook: WorkbookHandle,
        composition_id: str,
        content: WorkbookCompositionContent,
    ) -> WorkbookCompositionReference:
        """Compose one rectangular block of literal data OR formulas.

        This saves content without placing cells. Put headers, inputs and
        calculations in separate compositions/ranges when their formats differ.
        Never combine values with active formulas in one composition. Place each
        composition with put_range, then read_workbook before rendering.
        A column uses matrix values like [["January"], ["February"]]; row values
        run horizontally. Do not add placeholder cells where another range belongs.
        To edit content, reuse its composition_id. Resizing updates every placement
        and refuses collisions before saving.
        """
        if not composition_id.strip():
            raise ValueError('composition_id must be a non-empty string')
        validated = _CONTENT.validate_python(content)
        self._validate_content(validated)
        with self._workspace.lock:
            manifest = self._workspace.load(workbook)
            cast('dict[str, object]', manifest['compositions'])[composition_id] = validated
            self._refresh_range_bounds(manifest)
            self._workspace.write(workbook['workbook_id'], manifest)
        return {'workbook_id': workbook['workbook_id'], 'composition_id': composition_id}

    @toolify(permissions=PermissionSet(PermissionFlag.WRITE))
    def put_sheet(
        self,
        workbook: WorkbookHandle,
        sheet_id: str,
        name: str,
        position: WorkbookPosition,
        *,
        freeze_panes: str | None = None,
        print_settings: SheetPrintSettings | None = None,
    ) -> dict[str, object]:
        if (
            not sheet_id.strip()
            or not name.strip()
            or len(name) > 31
            or any(character in name for character in r'[]:*?/\\')
        ):
            raise ValueError('sheet_id and Excel sheet name must be valid')
        validated_position = _POSITION.validate_python(position)
        validated_print = (
            None if print_settings is None else _PRINT_SETTINGS.validate_python(print_settings)
        )
        if (
            validated_print is not None
            and not 50 <= validated_print.get('minimum_scale_percent', 70) <= 100
        ):
            raise ValueError('minimum_scale_percent must be between 50 and 100')
        if freeze_panes is not None:
            self._coordinate(freeze_panes)
        with self._workspace.lock:
            manifest = self._workspace.load(workbook)
            sheets = cast('list[dict[str, object]]', manifest['sheets'])
            if any(entry['name'] == name and entry['sheet_id'] != sheet_id for entry in sheets):
                raise ValueError(f"duplicate sheet name '{name}'")
            existing = next(
                (index for index, entry in enumerate(sheets) if entry['sheet_id'] == sheet_id),
                None,
            )
            created = existing is None
            ranges: list[dict[str, object]] = []
            if existing is not None:
                ranges = cast('list[dict[str, object]]', sheets[existing]['ranges'])
                if validated_print is None:
                    validated_print = cast(
                        'SheetPrintSettings | None', sheets[existing].get('print_settings')
                    )
                sheets.pop(existing)
            entry: dict[str, object] = {
                'sheet_id': sheet_id,
                'name': name,
                'position': validated_position,
                'freeze_panes': freeze_panes,
                'ranges': ranges,
                'print_settings': validated_print or {'mode': 'auto'},
            }
            sheets.insert(self._sheet_index(sheets, validated_position), entry)
            self._workspace.write(workbook['workbook_id'], manifest)
        return {'workbook': workbook, 'sheet_id': sheet_id, 'created': created}

    @requires_composition(
        'composition',
        upstream_tool='WORKBOOKS_compose_artifact',
        document_arg='workbook',
    )
    @toolify(permissions=PermissionSet(PermissionFlag.WRITE))
    def put_range(
        self,
        workbook: WorkbookHandle,
        range_id: str,
        sheet_id: str,
        start_cell: ExcelCell,
        composition: WorkbookCompositionReference,
        format: RangeFormat,
    ) -> dict[str, object]:
        """Place a composed rectangle, using one format for its whole range.

        Use separate non-overlapping ranges for labels, percentages, currencies
        and formulas. To revise a placement, reuse its exact existing range_id;
        adding another overlapping range does not restyle the original cells.
        """
        if not range_id.strip():
            raise ValueError('range_id must be a non-empty string')
        start_row, start_column = self._coordinate(start_cell)
        validated_format = _FORMAT.validate_python(format)
        self._validate_format(validated_format)
        with self._workspace.lock:
            manifest = self._workspace.load(workbook)
            sheet = self._sheet(manifest, sheet_id)
            content = cast('dict[str, WorkbookCompositionContent]', manifest['compositions'])[
                composition['composition_id']
            ]
            matrix = content_matrix(content)
            end_row = start_row + len(matrix) - 1
            end_column = start_column + len(matrix[0]) - 1
            if end_row > 1_048_576 or end_column > 16_384:
                raise ValueError('range exceeds XLSX worksheet bounds')
            if chart := validated_format.get('chart'):
                chart_series_columns(chart, start_column, end_column)
            ranges = cast('list[dict[str, object]]', sheet['ranges'])
            for entry in ranges:
                if entry['range_id'] == range_id:
                    continue
                if self._overlaps(
                    (start_row, start_column, end_row, end_column),
                    cast(
                        'tuple[int, int, int, int]',
                        tuple(cast('list[int]', entry['bounds'])),
                    ),
                ):
                    raise WorkbookRangeOverlapError(
                        f"range '{range_id}' would overlap '{entry['range_id']}'"
                    )
            replacement: dict[str, object] = {
                'range_id': range_id,
                'start_cell': start_cell,
                'composition': composition,
                'format': validated_format,
                'bounds': [start_row, start_column, end_row, end_column],
            }
            existing = next(
                (index for index, entry in enumerate(ranges) if entry['range_id'] == range_id),
                None,
            )
            created = existing is None
            if existing is None:
                ranges.append(replacement)
            else:
                ranges[existing] = replacement
            ranges.sort(
                key=lambda entry: (cast('list[int]', entry['bounds'])[0:2], entry['range_id'])
            )
            self._workspace.write(workbook['workbook_id'], manifest)
        return {'workbook': workbook, 'range_id': range_id, 'created': created}

    @toolify(permissions=PermissionSet(PermissionFlag.WRITE))
    def remove_range(self, workbook: WorkbookHandle, range_id: str) -> dict[str, object]:
        with self._workspace.lock:
            manifest = self._workspace.load(workbook)
            removed = False
            for sheet in cast('list[dict[str, object]]', manifest['sheets']):
                ranges = cast('list[dict[str, object]]', sheet['ranges'])
                retained = [entry for entry in ranges if entry['range_id'] != range_id]
                removed = removed or len(retained) != len(ranges)
                sheet['ranges'] = retained
            if removed:
                self._workspace.write(workbook['workbook_id'], manifest)
        return {'workbook': workbook, 'range_id': range_id, 'removed': removed}

    @toolify(permissions=PermissionSet(PermissionFlag.READ))
    @tool_output(_STATE_SCHEMA)
    def read_workbook(self, workbook: WorkbookHandle) -> WorkbookState:
        """Inspect placed content, charts and whether the output filename already exists.

        ``readback`` distinguishes formula compositions saved for later from formulas
        that have actual worksheet cells. Its lists are bounded: their counts cover
        the whole workbook and a true ``*_truncated`` flag identifies omitted detail.
        Before rendering verify the requested formulas are placed, chart series select
        the requested measures, and decide whether an existing destination may be
        overwritten. Reading does not change the workbook or output file.
        """
        with self._workspace.lock:
            manifest = self._workspace.load(workbook)
        compositions = cast('dict[str, WorkbookCompositionContent]', manifest['compositions'])
        states: list[WorkbookCompositionState] = [
            {'composition_id': key, 'content': value} for key, value in sorted(compositions.items())
        ]
        return {
            'workbook': workbook,
            'compositions': states,
            'sheets': cast('list[dict[str, object]]', manifest['sheets']),
            'readback': self._readback(manifest),
            'output_file_exists': self._workspace.safe_output_path(
                str(manifest['filename'])
            ).exists(),
        }

    @toolify(permissions=PermissionSet(PermissionFlag.WRITE))
    @tool_output(_GENERATED_SCHEMA)
    def render_workbook(
        self,
        workbook: WorkbookHandle,
        *,
        overwrite: bool = False,
        include_bytes: bool = False,
    ) -> GeneratedFile:
        """Render the placed workbook and return an immutable generated-file result.

        Check read_workbook first. If output_file_exists is true, overwrite=True
        is required to replace that local filename, and must be allowed by the
        user's task. Prior immutable artifact revisions remain unchanged.
        """
        with self._workspace.lock:
            manifest = self._workspace.load(workbook)
            if not cast('list[object]', manifest['sheets']):
                raise ValueError('workbook must contain at least one sheet')
            if not any(
                sheet['ranges'] for sheet in cast('list[dict[str, object]]', manifest['sheets'])
            ):
                raise ValueError(
                    'workbook needs placed content; compose and put_range before render'
                )
            raw = render_xlsx_bytes(manifest)
            target = self._workspace.write_output(
                str(manifest['filename']), raw, overwrite=overwrite
            )
            source_revision = snapshot_source(self._workspace.workspace_dir, manifest, raw)
            target = immutable_render_path(
                self._workspace.workspace_dir, source_revision, target.name, raw
            )
        return GeneratedFile(
            kind='generated_file',
            artifact_kind='spreadsheet',
            filename=str(manifest['filename']),
            path=str(target),
            mime_type=_XLSX_MIME,
            size_bytes=len(raw),
            sha256=hashlib.sha256(raw).hexdigest(),
            custody=source_custody(manifest),
            expected_base=expected_base(manifest),
            private_expected_base=expected_private_base(manifest),
            workspace={
                'source_revision': source_revision,
                'workbook_id': workbook['workbook_id'],
                'filename': workbook['filename'],
            },
            content_base64=base64.b64encode(raw).decode('ascii') if include_bytes else None,
        )

    @toolify(permissions=PermissionSet(PermissionFlag.WRITE))
    @tool_output(_GENERATED_SCHEMA)
    def create_xlsx(
        self,
        filename: str,
        document: OneShotWorkbook,
        *,
        overwrite: bool = False,
        include_bytes: bool = False,
    ) -> GeneratedFile:
        """Build and publish a COMPLETE workbook in one call.

        Include all headers, data, formulas and formatting in non-overlapping
        ranges. For incremental construction use create_workbook instead.
        Separate currency, percentage and label ranges; each range has one format.
        """
        validated = _ONE_SHOT.validate_python(document)
        target = self._workspace.safe_output_path(filename)
        if target.exists() and not overwrite:
            raise FileExistsError(str(target))
        workbook = self.create_workbook(filename)
        try:
            for sheet in validated['sheets']:
                self.put_sheet(
                    workbook,
                    sheet['sheet_id'],
                    sheet['name'],
                    {'anchor': 'end'},
                    freeze_panes=sheet.get('freeze_panes'),
                    print_settings=sheet.get('print_settings'),
                )
                for source in sheet['ranges']:
                    reference = self.compose_artifact(
                        workbook,
                        f'{sheet["sheet_id"]}-{source["range_id"]}-content',
                        source['content'],
                    )
                    self.put_range(
                        workbook,
                        source['range_id'],
                        sheet['sheet_id'],
                        source['start_cell'],
                        reference,
                        source.get('format', {'style': 'body'}),
                    )
            return self.render_workbook(workbook, overwrite=overwrite, include_bytes=include_bytes)
        except Exception:
            self._workspace.discard(workbook)
            raise

    def validate_composition_reference(
        self,
        doc: object,
        reference: object,
        argument_name: str,
    ) -> None:
        required = 'WORKBOOKS_compose_artifact'
        if not isinstance(doc, Mapping) or not isinstance(reference, Mapping):
            raise CompositionDependencyError(
                f"argument '{argument_name}' requires a reference returned by {required}"
            )
        doc_map = cast('Mapping[object, object]', doc)
        ref_map = cast('Mapping[object, object]', reference)
        composition_id = ref_map.get('composition_id')
        if ref_map.get('workbook_id') != doc_map.get('workbook_id') or not isinstance(
            composition_id, str
        ):
            raise CompositionDependencyError(
                f"argument '{argument_name}' requires a reference returned by {required}"
            )
        manifest = self._workspace.load(cast('WorkbookHandle', doc))
        if composition_id not in cast('dict[str, object]', manifest['compositions']):
            raise CompositionDependencyError(
                f"argument '{argument_name}' references unknown composition '{composition_id}'; "
                f'call {required} first'
            )

    @staticmethod
    def _refresh_range_bounds(manifest: dict[str, object]) -> None:
        compositions = cast('dict[str, WorkbookCompositionContent]', manifest['compositions'])
        for sheet in cast('list[dict[str, object]]', manifest['sheets']):
            placed: list[tuple[int, int, int, int]] = []
            for entry in cast('list[dict[str, object]]', sheet['ranges']):
                reference = cast('WorkbookCompositionReference', entry['composition'])
                matrix = content_matrix(compositions[reference['composition_id']])
                row, column = WorkbooksToolSet._coordinate(entry['start_cell'])
                bounds = (row, column, row + len(matrix) - 1, column + len(matrix[0]) - 1)
                if chart := cast('RangeFormat', entry.get('format', {})).get('chart'):
                    chart_series_columns(chart, column, bounds[3])
                if bounds[2] > 1_048_576 or bounds[3] > 16_384:
                    raise ValueError('range exceeds XLSX worksheet bounds')
                if any(WorkbooksToolSet._overlaps(bounds, prior) for prior in placed):
                    raise WorkbookRangeOverlapError(
                        'composition resize would overlap an existing range'
                    )
                entry['bounds'] = list(bounds)
                placed.append(bounds)

    @staticmethod
    def _readback(manifest: dict[str, object]) -> WorkbookReadback:
        """Summarize actual placements without duplicating every composed workbook cell."""
        compositions = cast('dict[str, WorkbookCompositionContent]', manifest['compositions'])
        sheets = cast('list[dict[str, object]]', manifest['sheets'])
        placed_ranges: list[WorkbookPlacedRangeReadback] = []
        formula_addresses: list[WorkbookFormulaAddress] = []
        sheet_summaries: list[WorkbookSheetReadback] = []
        placed_composition_ids: set[str] = set()
        total_range_count = 0
        total_formula_count = 0

        for sheet in sheets:
            ranges = cast('list[dict[str, object]]', sheet['ranges'])
            sheet_range_count = 0
            sheet_cell_count = 0
            sheet_formula_count = 0
            minimum_row: int | None = None
            minimum_column: int | None = None
            maximum_row = 0
            maximum_column = 0
            for entry in ranges:
                reference = cast('WorkbookCompositionReference', entry['composition'])
                composition_id = reference['composition_id']
                content = compositions[composition_id]
                matrix = content_matrix(content)
                start_row, start_column = WorkbooksToolSet._coordinate(entry['start_cell'])
                row_count = len(matrix)
                column_count = len(matrix[0])
                end_row = start_row + row_count - 1
                end_column = start_column + column_count - 1
                formula_count = WorkbooksToolSet._formula_count(content)
                cell_count = row_count * column_count
                range_id = cast('str', entry['range_id'])
                sheet_id = cast('str', sheet['sheet_id'])

                placed_composition_ids.add(composition_id)
                total_range_count += 1
                total_formula_count += formula_count
                sheet_range_count += 1
                sheet_cell_count += cell_count
                sheet_formula_count += formula_count
                minimum_row = start_row if minimum_row is None else min(minimum_row, start_row)
                minimum_column = (
                    start_column if minimum_column is None else min(minimum_column, start_column)
                )
                maximum_row = max(maximum_row, end_row)
                maximum_column = max(maximum_column, end_column)

                if len(placed_ranges) < _READBACK_MAX_RANGES:
                    placed_ranges.append(
                        {
                            'sheet_id': sheet_id,
                            'sheet_name': cast('str', sheet['name']),
                            'range_id': range_id,
                            'composition_id': composition_id,
                            'start_cell': cast('ExcelCell', entry['start_cell']),
                            'end_cell': f'{get_column_letter(end_column)}{end_row}',
                            'row_count': row_count,
                            'column_count': column_count,
                            'covered_cell_count': cell_count,
                            'formula_count': formula_count,
                        }
                    )
                for row_offset, row in enumerate(WorkbooksToolSet._formula_matrix(content)):
                    for column_offset, formula in enumerate(row):
                        if (
                            formula is not None
                            and len(formula_addresses) < _READBACK_MAX_FORMULA_ADDRESSES
                        ):
                            formula_addresses.append(
                                {
                                    'sheet_id': sheet_id,
                                    'range_id': range_id,
                                    'address': (
                                        f'{get_column_letter(start_column + column_offset)}'
                                        f'{start_row + row_offset}'
                                    ),
                                }
                            )

            if len(sheet_summaries) < _READBACK_MAX_SHEETS:
                sheet_summaries.append(
                    {
                        'sheet_id': cast('str', sheet['sheet_id']),
                        'sheet_name': cast('str', sheet['name']),
                        'placed_range_count': sheet_range_count,
                        'covered_cell_count': sheet_cell_count,
                        'formula_count': sheet_formula_count,
                        'row_count': 0 if minimum_row is None else maximum_row - minimum_row + 1,
                        'column_count': (
                            0 if minimum_column is None else maximum_column - minimum_column + 1
                        ),
                    }
                )

        unplaced_composition_ids = sorted(set(compositions) - placed_composition_ids)
        unplaced_compositions: list[WorkbookUnplacedCompositionReadback] = []
        for composition_id in unplaced_composition_ids[:_READBACK_MAX_UNPLACED_COMPOSITIONS]:
            content = compositions[composition_id]
            matrix = content_matrix(content)
            unplaced_compositions.append(
                {
                    'composition_id': composition_id,
                    'kind': content['kind'],
                    'row_count': len(matrix),
                    'column_count': len(matrix[0]),
                    'formula_count': WorkbooksToolSet._formula_count(content),
                }
            )

        return {
            'sheet_count': len(sheets),
            'sheets': sheet_summaries,
            'sheets_truncated': len(sheets) > len(sheet_summaries),
            'placed_range_count': total_range_count,
            'placed_ranges': placed_ranges,
            'placed_ranges_truncated': total_range_count > len(placed_ranges),
            'placed_formula_count': total_formula_count,
            'placed_formula_addresses': formula_addresses,
            'placed_formula_addresses_truncated': total_formula_count > len(formula_addresses),
            'unplaced_composition_count': len(unplaced_composition_ids),
            'unplaced_compositions': unplaced_compositions,
            'unplaced_compositions_truncated': (
                len(unplaced_composition_ids) > len(unplaced_compositions)
            ),
        }

    @staticmethod
    def _formula_matrix(content: WorkbookCompositionContent) -> list[list[str | None]]:
        return content.get('formulas') or []

    @staticmethod
    def _formula_count(content: WorkbookCompositionContent) -> int:
        return sum(
            formula is not None
            for row in WorkbooksToolSet._formula_matrix(content)
            for formula in row
        )

    @staticmethod
    def _validate_content(content: WorkbookCompositionContent) -> None:
        if content['kind'] == 'formulas' and (
            content.get('value') not in (None, '')
            or any(
                cell not in (None, '')
                for item in (content.get('values') or [])
                for cell in (item if isinstance(item, list) else [item])
            )
        ):
            raise _content_error(
                'workbook_literals_and_formulas_require_separate_compositions',
                'put literal data in a separate scalar, row or matrix composition',
                'values',
            )
        if content['kind'] != 'formulas' and any(
            isinstance(formula, str) and formula.startswith('=')
            for row in (content.get('formulas') or [])
            for formula in row
        ):
            raise _content_error(
                'workbook_literals_and_formulas_require_separate_compositions',
                'put active formulas in a separate kind=formulas composition',
                'formulas',
            )
        matrix = content_matrix(content)
        field = 'formulas' if content['kind'] == 'formulas' else 'values'
        if any(not isinstance(row, list) for row in cast('list[object]', matrix)):
            raise _content_error(
                'workbook_matrix_requires_rows',
                'matrix content needs rows, for example [[100], [120]]',
                field,
            )
        if any(isinstance(cell, list) for row in matrix for cell in row):
            raise _content_error(
                'workbook_row_requires_flat_values',
                'row content needs flat values; use matrix for multiple rows',
                field,
            )
        if not matrix or not matrix[0] or any(len(row) != len(matrix[0]) for row in matrix):
            raise _content_error(
                'workbook_matrix_requires_equal_length_rows',
                'workbook compositions need non-empty rows of equal length in a rectangular range',
                field,
            )
        if content['kind'] == 'formulas':
            for row_index, row in enumerate(matrix):
                for column_index, formula in enumerate(row):
                    if formula is not None:
                        rule = 'workbook_formula_requires_supported_equals_expression_or_null'
                        message = 'formula must be a supported expression starting with =, or null'
                        if not isinstance(formula, str):
                            raise _content_error(rule, message, field, row_index, column_index)
                        try:
                            validate_formula(formula)
                        except ValueError:
                            raise _content_error(
                                rule, message, field, row_index, column_index
                            ) from None

    @staticmethod
    def _validate_format(formatting: RangeFormat) -> None:
        number_format = formatting.get('number_format')
        if number_format is not None and number_format not in _NUMBER_FORMATS:
            raise ValueError('number_format is outside the supported allowlist')
        merge = formatting.get('merge')
        if merge is not None:
            first, last = merge.split(':')
            first_row, first_column = WorkbooksToolSet._coordinate(first)
            last_row, last_column = WorkbooksToolSet._coordinate(last)
            if first_row > last_row or first_column > last_column:
                raise ValueError('merge must end at or after its starting cell')
        chart = formatting.get('chart')
        if chart is not None:
            WorkbooksToolSet._coordinate(chart['anchor'])
            source = chart.get('source_range')
            if source is not None:
                first, last = source.split(':')
                first_row, first_column = WorkbooksToolSet._coordinate(first)
                last_row, last_column = WorkbooksToolSet._coordinate(last)
                if first_row >= last_row or first_column >= last_column:
                    raise ValueError('chart source needs headers, categories and numeric data')

    @staticmethod
    def _coordinate(value: object) -> tuple[int, int]:
        if (
            not isinstance(value, str)
            or re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]{0,6}', value) is None
        ):
            raise ValueError(f"invalid worksheet coordinate '{value}'")
        try:
            row, column = coordinate_to_tuple(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"invalid worksheet coordinate '{value}'") from exc
        if row > 1_048_576 or column > 16_384:
            raise ValueError('worksheet coordinate exceeds XLSX bounds')
        return row, column

    @staticmethod
    def _sheet(manifest: dict[str, object], sheet_id: str) -> dict[str, object]:
        for sheet in cast('list[dict[str, object]]', manifest['sheets']):
            if sheet['sheet_id'] == sheet_id:
                return sheet
        raise ValueError(f"unknown sheet_id '{sheet_id}'")

    @staticmethod
    def _sheet_index(sheets: list[dict[str, object]], position: WorkbookPosition) -> int:
        if position['anchor'] == 'start':
            return 0
        if position['anchor'] == 'end':
            return len(sheets)
        anchor_id = position.get('sheet_id')
        if not anchor_id:
            raise ValueError(f"position anchor '{position['anchor']}' requires sheet_id")
        for index, sheet in enumerate(sheets):
            if sheet['sheet_id'] == anchor_id:
                return index if position['anchor'] == 'before' else index + 1
        raise ValueError(f"position references unknown sheet_id '{anchor_id}'")

    @staticmethod
    def _overlaps(left: tuple[int, int, int, int], right: tuple[int, int, int, int]) -> bool:
        return not (
            left[2] < right[0] or right[2] < left[0] or left[3] < right[1] or right[3] < left[1]
        )
