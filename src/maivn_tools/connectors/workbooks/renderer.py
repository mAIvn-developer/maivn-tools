"""Deterministic XLSX rendering from a persisted workbook manifest."""

# pyright: strict

from __future__ import annotations

import io
import math
import re
import textwrap
from datetime import datetime, timezone
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.properties import PageSetupProperties
from pydantic import ValidationError
from pydantic_core import PydanticCustomError

from ..artifacts import canonicalize_zip
from .cached_results import cache_results
from .models import RangeChart, RangeFormat, SheetPrintSettings, WorkbookCompositionContent

_FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


def chart_series_columns(chart: RangeChart, first_column: int, last_column: int) -> list[int]:
    """Resolve explicit series against the source, including separately placed ranges."""
    if source := chart.get('source_range'):
        first_column, _, last_column, _ = cast(
            'tuple[int, int, int, int]', range_boundaries(source)
        )
    columns = chart.get('series_columns')
    if columns is None:
        return list(range(first_column + 1, last_column + 1))
    if len(columns) != len(set(columns)) or any(
        column <= first_column or column > last_column for column in columns
    ):
        raise ValueError('chart series columns must be unique measure columns inside the source')
    return columns


def content_matrix(content: WorkbookCompositionContent) -> list[list[object]]:
    if content['kind'] == 'scalar':
        return [[content.get('value')]]
    if content['kind'] == 'row':
        return [cast('list[object]', content.get('values') or [])]
    if content['kind'] == 'matrix':
        return cast('list[list[object]]', content.get('values') or [])
    return cast('list[list[object]]', content.get('formulas') or [])


def render_xlsx_bytes(manifest: dict[str, object]) -> bytes:
    workbook = Workbook()
    active = workbook.active
    if active is not None:
        workbook.remove(active)
    properties = workbook.properties
    properties.creator = ''
    properties.lastModifiedBy = ''
    properties.title = ''
    properties.subject = ''
    properties.description = ''
    properties.keywords = ''
    properties.category = ''
    properties.created = _FIXED_TIME
    properties.modified = _FIXED_TIME
    properties.lastPrinted = None
    workbook.calculation = CalcProperties(
        calcMode='auto',
        fullCalcOnLoad=True,
        forceFullCalc=True,
    )
    compositions = cast('dict[str, WorkbookCompositionContent]', manifest['compositions'])
    sheets = cast('list[dict[str, object]]', manifest['sheets'])
    for sheet_entry in sheets:
        sheet = workbook.create_sheet(cast('str', sheet_entry['name']))
        sheet.sheet_view.showGridLines = False
        freeze = sheet_entry.get('freeze_panes')
        if isinstance(freeze, str):
            sheet.freeze_panes = freeze
        for range_entry in cast('list[dict[str, object]]', sheet_entry['ranges']):
            reference = cast('dict[str, str]', range_entry['composition'])
            content = compositions[reference['composition_id']]
            matrix = content_matrix(content)
            start_row, start_column = coordinate_to_tuple(cast('str', range_entry['start_cell']))
            for row_offset, row in enumerate(matrix):
                for column_offset, value in enumerate(row):
                    cell = sheet.cell(start_row + row_offset, start_column + column_offset, value)
                    if isinstance(value, str) and content['kind'] != 'formulas':
                        # Input data is literal even when it starts with an equals sign.
                        cell.data_type = 's'
            end_row = start_row + len(matrix) - 1
            end_column = start_column + len(matrix[0]) - 1
            formatting = cast('RangeFormat', range_entry.get('format', {}))
            _apply_format(sheet, start_row, start_column, end_row, end_column, formatting)
            merge = formatting.get('merge')
            if merge:
                sheet.merge_cells(merge)
            chart_spec = formatting.get('chart')
            if chart_spec is not None:
                chart = BarChart() if chart_spec['type'] == 'bar' else LineChart()
                chart.title = chart_spec['title']
                cast('Any', chart.title).overlay = False
                chart.style = 13
                chart.height = 7
                chart.width = 12
                source = chart_spec.get('source_range')
                source_bounds = (
                    range_boundaries(source)
                    if source
                    else (start_column, start_row, end_column, end_row)
                )
                first_col, first_row, last_col, last_row = cast(
                    'tuple[int, int, int, int]', source_bounds
                )
                series_columns = chart_series_columns(chart_spec, first_col, last_col)
                for column in series_columns:
                    data = Reference(
                        sheet,
                        min_col=column,
                        max_col=column,
                        min_row=first_row,
                        max_row=last_row,
                    )
                    cast('Any', chart).add_data(data, titles_from_data=True)
                categories = Reference(
                    sheet,
                    min_col=first_col,
                    min_row=first_row + 1,
                    max_row=last_row,
                )
                cast('Any', chart).set_categories(categories)
                chart.x_axis.delete = False
                chart.y_axis.delete = False
                chart.x_axis.tickLblPos = 'nextTo'
                chart.y_axis.tickLblPos = 'nextTo'
                chart.y_axis.numFmt = 'General'
                for index, series in enumerate(cast('list[Any]', chart.series)):
                    color = ('24496B', '0F766E', 'B45309', '7C3AED')[index % 4]
                    series.graphicalProperties.solidFill = color
                    series.graphicalProperties.line.solidFill = color
                if len(series_columns) == 1:
                    chart.legend = None
                elif chart.legend is not None:
                    chart.legend.position = 'b'
                    chart.legend.overlay = False
                sheet.add_chart(chart, chart_spec['anchor'])
        _size_content(sheet, sheet_entry, compositions)
        _apply_print_settings(sheet, sheet_entry)
    raw = io.BytesIO()
    workbook.save(raw)
    return canonicalize_zip(
        cache_results(raw.getvalue(), workbook), member_transform=_canonical_member
    )


def _size_content(
    sheet: Worksheet,
    entry: dict[str, object],
    compositions: dict[str, WorkbookCompositionContent],
) -> None:
    """Measure every placed range before sizing so later ranges cannot shrink headers."""
    widths: dict[int, float] = {}
    cells: set[tuple[int, int]] = set()
    for placed in cast('list[dict[str, object]]', entry['ranges']):
        reference = cast('dict[str, str]', placed['composition'])
        matrix = content_matrix(compositions[reference['composition_id']])
        start_row, start_column = coordinate_to_tuple(cast('str', placed['start_cell']))
        end_row, end_column = start_row + len(matrix) - 1, start_column + len(matrix[0]) - 1
        formatting = cast('RangeFormat', placed.get('format', {}))
        if formatting.get('style') == 'header_row' and sheet.print_title_rows is None:
            sheet.print_title_rows = f'{start_row}:{start_row}'
        for row in sheet.iter_rows(
            min_row=start_row, max_row=end_row, min_col=start_column, max_col=end_column
        ):
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value)
                # Formula source length says nothing about displayed result length.
                length = (
                    14 if cell.data_type == 'f' else max(len(line) for line in text.split('\n'))
                )
                widths[cell.column] = max(widths.get(cell.column, 10), min(40, length + 2))
                cells.add((cell.row, cell.column))
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    heights: dict[int, float] = {}
    for row, column in cells:
        cell = sheet.cell(row, column)
        width = widths[column]
        if cell.data_type == 'f' or not isinstance(cell.value, str):
            lines = 1
        else:
            lines = sum(
                max(1, len(textwrap.wrap(line, width=max(1, math.floor(width - 2)))))
                for line in cell.value.split('\n')
            )
        height = min(409, max(22, lines * 15 + 8))
        heights[row] = max(heights.get(row, 0), height)
    for row, height in heights.items():
        sheet.row_dimensions[row].height = height


class WorkbookPrintLayoutError(ValidationError):
    """A workbook-state repair can make the identical render request valid."""

    sdk_retryable = True


def _apply_print_settings(sheet: Worksheet, entry: dict[str, object]) -> None:
    """Fit compact sheets including charts; keep large sheets at readable scale."""
    settings = cast('SheetPrintSettings', entry.get('print_settings', {'mode': 'auto'}))
    orientation = settings.get('orientation', 'portrait')
    printable_inches = (11 if orientation == 'landscape' else 8.5) - 1.5
    chart_columns: list[int] = []
    for item in cast('list[dict[str, object]]', entry['ranges']):
        chart = cast('RangeFormat', item.get('format', {})).get('chart')
        if chart is not None:
            chart_columns.append(coordinate_to_tuple(chart['anchor'])[1])
    max_column = max(sheet.max_column, max(chart_columns, default=1))
    # Excel character-width estimates, including gaps before chart anchors.
    offsets = [0.0]
    for column in range(1, max_column + 1):
        dimension = sheet.column_dimensions.get(get_column_letter(column))
        width = dimension.width if dimension is not None else 13.0
        offsets.append(offsets[-1] + (width * 7 + 5) / 96)
    right_edge = offsets[sheet.max_column]
    for column in chart_columns:
        right_edge = max(right_edge, offsets[column - 1] + 12 / 2.54)
    estimated_scale = min(100.0, printable_inches / max(right_edge, 0.01) * 100)
    readable = estimated_scale >= settings.get('minimum_scale_percent', 70)
    mode = settings['mode']
    if mode == 'fit_width' and not readable:
        raise WorkbookPrintLayoutError.from_exception_data(
            'WorkbookPrintSettings',
            [
                {
                    'type': PydanticCustomError(
                        'workbook_print_requires_landscape_or_actual_size',
                        'sheet cannot fit at minimum print scale; use landscape or actual_size',
                    ),
                    'loc': ('print_settings',),
                    'input': None,
                }
            ],
            hide_input=True,
        )
    fit = mode != 'actual_size' and readable
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=fit)
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1 if fit else 0
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.scale = None if fit else 100


def _canonical_member(name: str, content: bytes) -> bytes:
    if name != 'docProps/core.xml':
        return content
    canonical, replacements = re.subn(
        rb'(<dcterms:modified\b[^>]*>).*?(</dcterms:modified>)',
        rb'\g<1>2000-01-01T00:00:00Z\g<2>',
        content,
        count=1,
    )
    if replacements != 1:
        raise ValueError('XLSX core properties must contain exactly one modified timestamp')
    return canonical


def _apply_format(
    sheet: Worksheet,
    start_row: int,
    start_column: int,
    end_row: int,
    end_column: int,
    formatting: RangeFormat,
) -> None:
    style = formatting.get('style') or 'body'
    number_format = formatting.get('number_format') or 'General'
    for row in sheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_column,
        max_col=end_column,
    ):
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.number_format = number_format
            if style == 'header_row' and cell.row == start_row:
                cell.font = Font(name='Aptos', size=11, bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='24496B')
            elif style == 'total':
                cell.font = Font(name='Aptos', size=11, bold=True, color='183153')
                cell.fill = PatternFill('solid', fgColor='DCE8F2')
            elif style == 'input':
                cell.font = Font(name='Aptos', size=11, color='0000FF')
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
            elif style == 'accent':
                cell.font = Font(name='Aptos', size=11, bold=True, color='183153')
                cell.fill = PatternFill('solid', fgColor='E2F0D9')
            else:
                cell.font = Font(name='Aptos', size=11, color='000000')
