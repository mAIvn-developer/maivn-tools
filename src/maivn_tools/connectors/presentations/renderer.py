"""Deterministic PPTX rendering from a persisted presentation manifest."""

# pyright: strict

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any, cast

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from ..artifacts import canonicalize_zip
from ..artifacts.images import ImageReference, contain_size, image_bytes
from .models import LIST_INDENT_INCHES, PresentationCompositionContent, PresentationElement

_FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


def render_pptx_bytes(manifest: dict[str, object]) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    properties = presentation.core_properties
    properties.title = ''
    properties.subject = ''
    properties.author = ''
    properties.keywords = ''
    properties.comments = ''
    properties.category = ''
    properties.last_modified_by = ''
    properties.created = _FIXED_TIME
    properties.modified = _FIXED_TIME
    properties.revision = 1
    compositions = cast('dict[str, PresentationCompositionContent]', manifest['compositions'])
    for slide_entry in cast('list[dict[str, object]]', manifest['slides']):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(248, 250, 252)
        for element_entry in cast('list[dict[str, object]]', slide_entry['elements']):
            element = cast('PresentationElement', element_entry['element'])
            content = compositions[element['composition']['composition_id']]
            left = Inches(element['x'])
            top = Inches(element['y'])
            width = Inches(element['width'])
            height = Inches(element['height'])
            if element['type'] == 'image':
                image = cast('ImageReference', content.get('image'))
                raw_image = image_bytes(manifest, image)
                fit_width, fit_height = contain_size(raw_image, element['width'], element['height'])
                picture = slide.shapes.add_picture(
                    io.BytesIO(raw_image),
                    Inches(element['x'] + (element['width'] - fit_width) / 2),
                    Inches(element['y'] + (element['height'] - fit_height) / 2),
                    Inches(fit_width),
                    Inches(fit_height),
                )
                cast('Any', picture.element).nvPicPr.cNvPr.set('descr', image['alt_text'])
            elif element['type'] in {'title', 'text', 'list'}:
                box = slide.shapes.add_textbox(left, top, width, height)
                frame = box.text_frame
                frame.clear()
                frame.word_wrap = True
                frame.margin_left = Inches(0.08)
                frame.margin_right = Inches(0.08)
                frame.margin_top = Inches(0.04)
                frame.margin_bottom = Inches(0.04)
                frame.vertical_anchor = MSO_ANCHOR.TOP
                if element['type'] == 'list':
                    items = cast('list[str]', content.get('items'))
                    for index, item in enumerate(items):
                        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                        paragraph.text = item
                        paragraph.level = 0
                        # python-pptx has no public API for native bullet properties.
                        properties = cast('Any', paragraph)._p.get_or_add_pPr()  # noqa: SLF001
                        properties.set('marL', str(Inches(LIST_INDENT_INCHES)))
                        properties.set('indent', str(-Inches(0.18)))
                        bullet = cast('Any', OxmlElement('a:buChar'))
                        bullet.set('char', '•')
                        properties.append(bullet)
                        paragraph.font.size = Pt(20)
                        paragraph.font.name = 'Aptos'
                        paragraph.font.color.rgb = RGBColor(36, 73, 107)
                        paragraph.space_after = Pt(8)
                else:
                    paragraph = frame.paragraphs[0]
                    paragraph.text = cast('str', content.get('text'))
                    paragraph.font.name = 'Aptos Display' if element['type'] == 'title' else 'Aptos'
                    paragraph.font.size = Pt(36 if element['type'] == 'title' else 20)
                    paragraph.font.bold = element['type'] == 'title'
                    paragraph.font.color.rgb = _text_color(element.get('style', 'default'))
                    paragraph.alignment = PP_ALIGN.LEFT
                for paragraph in frame.paragraphs:
                    paragraph.font.size = Pt(
                        element.get('font_size', 36 if element['type'] == 'title' else 20)
                    )
                    paragraph.font.name = element.get(
                        'font_family', 'Aptos Display' if element['type'] == 'title' else 'Aptos'
                    )
                    paragraph.alignment = {
                        'left': PP_ALIGN.LEFT,
                        'center': PP_ALIGN.CENTER,
                        'right': PP_ALIGN.RIGHT,
                    }[element.get('alignment', 'left')]
            elif element['type'] == 'table':
                rows = cast('list[list[str]]', content.get('rows'))
                shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, height)
                table = shape.table
                for row_index, row in enumerate(rows):
                    for column_index, value in enumerate(row):
                        cell = table.cell(row_index, column_index)
                        cell.text = value
                        cell.margin_left = Inches(0.06)
                        cell.margin_right = Inches(0.06)
                        cell.margin_top = Inches(0.04)
                        cell.margin_bottom = Inches(0.04)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        cell.text_frame.word_wrap = True
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.name = element.get('font_family', 'Aptos')
                            paragraph.font.size = Pt(element.get('font_size', 14))
                            paragraph.font.bold = row_index == 0
                            paragraph.alignment = {
                                'left': PP_ALIGN.LEFT,
                                'center': PP_ALIGN.CENTER,
                                'right': PP_ALIGN.RIGHT,
                            }[element.get('alignment', 'left')]
                            paragraph.font.color.rgb = (
                                RGBColor(255, 255, 255) if row_index == 0 else RGBColor(24, 49, 83)
                            )
                        fill = cell.fill
                        fill.solid()
                        fill.fore_color.rgb = (
                            RGBColor(36, 73, 107) if row_index == 0 else RGBColor(226, 232, 240)
                        )
            else:
                data = ChartData()
                categories = cast('list[str]', content.get('categories'))
                series_items = cast('list[dict[str, object]]', content.get('series'))
                data.categories = categories
                for series in series_items:
                    cast('Any', data).add_series(
                        cast('str', series['name']),
                        tuple(cast('list[float]', series['values'])),
                    )
                chart = cast(
                    'Any',
                    slide.shapes.add_chart(
                        XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, data
                    ),
                ).chart
                # Even one series can carry the measure and units in its name.
                # Otherwise PowerPoint substitutes category names in its legend.
                chart.plots[0].vary_by_categories = False
                chart.has_legend = True
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
                chart.legend.font.name = element.get('font_family', 'Aptos')
                chart.legend.font.size = Pt(element.get('font_size', 18))
                for axis in (chart.category_axis, chart.value_axis):
                    axis.tick_labels.font.name = element.get('font_family', 'Aptos')
                    axis.tick_labels.font.size = Pt(element.get('font_size', 18))
                chart.has_title = False
                chart.value_axis.has_major_gridlines = True
    raw = io.BytesIO()
    presentation.save(raw)
    return canonicalize_zip(raw.getvalue(), member_transform=_canonical_member)


def _canonical_member(name: str, content: bytes) -> bytes:
    if name.startswith('ppt/charts/') and name.endswith('.xml'):
        return _canonical_chart_member(content)
    if not name.startswith('ppt/embeddings/') or not name.endswith('.xlsx'):
        return content
    return canonicalize_zip(content, member_transform=_canonical_embedded_workbook_member)


def _canonical_chart_member(content: bytes) -> bytes:
    # OOXML requires unsigned IDs, but PowerPoint rejects values above Int32.
    # IDs are chart-local references: renumber every identity together so both
    # axis definitions and cross-references remain distinct and deterministic.
    axis_ids: dict[int, int] = {}

    def canonical_axis_id(match: re.Match[bytes]) -> bytes:
        original = int(match.group(2))
        value = axis_ids.setdefault(original, len(axis_ids) + 1)
        return match.group(1) + str(value).encode('ascii') + match.group(3)

    return re.sub(
        rb'(<c:(?:axId|crossAx)\s+val=")(-?\d+)(")',
        canonical_axis_id,
        content,
    )


def _canonical_embedded_workbook_member(name: str, content: bytes) -> bytes:
    if name != 'docProps/core.xml':
        return content
    for field in (b'created', b'modified'):
        content, replacements = re.subn(
            rb'(<dcterms:' + field + rb'\b[^>]*>).*?(</dcterms:' + field + rb'>)',
            rb'\g<1>2000-01-01T00:00:00Z\g<2>',
            content,
            count=1,
        )
        if replacements != 1:
            raise ValueError(f'embedded chart workbook must contain one {field.decode()} timestamp')
    return content


def _text_color(style: str) -> RGBColor:
    if style == 'accent':
        return RGBColor(0, 112, 173)
    if style == 'muted':
        return RGBColor(100, 116, 139)
    return RGBColor(24, 49, 83)
