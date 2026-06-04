"""Optional, dependency-gated text extractors.

This module never imports its optional dependencies at module load. Each
extractor checks for its driver only when ``extract`` runs, so importing
``maivn_tools.files.extra_extractors`` is safe regardless of which extras
the host installed.

Available extractors:

* :class:`HtmlTextExtractor` — pure stdlib (no extras required).
* :class:`PdfTextExtractor` — requires the ``pdf`` extra (``pypdf``).
* :class:`DocxTextExtractor` — requires the ``docx`` extra (``python-docx``).

Call :func:`register_default_extractors` once at process start to plug them
into the global extractor registry.
"""
# pyright: strict

from __future__ import annotations

import io
from html.parser import HTMLParser

from .extractors import ExtractionResult, TextExtractor, register_extractor

# MARK: HTML extractor


class _PlainTextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._chunks: list[str] = []
        self._skip_depth = 0
        self._skip_tags = {"script", "style", "head"}

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._skip_tags:
            self._skip_depth += 1
        elif tag in {"br", "p", "li", "div", "tr"}:
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._skip_tags and self._skip_depth > 0:
            self._skip_depth -= 1
        elif tag in {"p", "li", "div", "tr"}:
            self._chunks.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0:
            self._chunks.append(data)

    @property
    def text(self) -> str:
        joined = "".join(self._chunks)
        # Collapse runs of blank lines and strip leading/trailing whitespace.
        lines = [line.strip() for line in joined.splitlines()]
        cleaned: list[str] = []
        previous_blank = False
        for line in lines:
            if not line:
                if previous_blank:
                    continue
                previous_blank = True
            else:
                previous_blank = False
            cleaned.append(line)
        return "\n".join(cleaned).strip()


class HtmlTextExtractor(TextExtractor):
    """Pure-stdlib HTML to plain-text extractor.

    The extractor strips ``<script>``, ``<style>``, and ``<head>`` blocks,
    collapses block-level tags into newlines, and decodes HTML entities. It
    is intentionally minimal; callers that need DOM-level extraction should
    plug in a richer parser (e.g. ``readability-lxml``).
    """

    mime_types = ("text/html", "application/xhtml+xml")

    def extract(
        self,
        data: bytes,
        *,
        mime_type: str | None = None,
        filename: str | None = None,
    ) -> ExtractionResult:
        text_in = data.decode("utf-8", errors="replace") if data else ""
        parser = _PlainTextHTMLParser()
        parser.feed(text_in)
        parser.close()
        return ExtractionResult(
            text=parser.text,
            mime_type="text/plain",
            metadata={"source_mime": mime_type or "text/html"},
        )


# MARK: PDF extractor


class PdfTextExtractor(TextExtractor):
    """PDF extractor backed by ``pypdf``.

    The dependency is loaded lazily inside :meth:`extract` so the module
    remains import-safe even when the ``pdf`` extra is not installed.
    """

    mime_types = ("application/pdf",)

    def extract(
        self,
        data: bytes,
        *,
        mime_type: str | None = None,
        filename: str | None = None,
    ) -> ExtractionResult:
        try:
            from pypdf import PdfReader  # type: ignore[import-not-found]
        except ImportError as exc:  # pragma: no cover - depends on env
            raise RuntimeError(
                "PdfTextExtractor requires the 'pdf' extra. "
                "Install with: pip install maivn-tools[pdf]"
            ) from exc
        reader = PdfReader(io.BytesIO(data))
        pages: list[str] = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return ExtractionResult(
            text="\n\n".join(pages).strip(),
            mime_type="text/plain",
            metadata={"page_count": len(reader.pages), "source_mime": "application/pdf"},
        )


# MARK: DOCX extractor


class DocxTextExtractor(TextExtractor):
    """DOCX extractor backed by ``python-docx``."""

    mime_types = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",)

    def extract(
        self,
        data: bytes,
        *,
        mime_type: str | None = None,
        filename: str | None = None,
    ) -> ExtractionResult:
        try:
            from docx import Document  # type: ignore[import-not-found]
        except ImportError as exc:  # pragma: no cover - depends on env
            raise RuntimeError(
                "DocxTextExtractor requires the 'docx' extra. "
                "Install with: pip install maivn-tools[docx]"
            ) from exc
        document = Document(io.BytesIO(data))
        paragraphs = [para.text for para in document.paragraphs if para.text]
        return ExtractionResult(
            text="\n".join(paragraphs).strip(),
            mime_type="text/plain",
            metadata={
                "paragraph_count": len(document.paragraphs),
                "source_mime": mime_type or "docx",
            },
        )


# MARK: Registration


def register_default_extractors() -> None:
    """Register the bundled HTML / PDF / DOCX extractors globally.

    Each extractor checks for its driver lazily inside ``extract``, so
    registering all three is always safe; calls to extractors whose drivers
    are missing will raise a clear ``RuntimeError`` only at use time.
    """
    register_extractor(HtmlTextExtractor())
    register_extractor(PdfTextExtractor())
    register_extractor(DocxTextExtractor())


__all__ = [
    "DocxTextExtractor",
    "HtmlTextExtractor",
    "PdfTextExtractor",
    "register_default_extractors",
]
