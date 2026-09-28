"""Repeatable forecast quality fixture using only the public workbook toolset."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import TYPE_CHECKING, Any, cast
from zipfile import ZipFile

from openpyxl import load_workbook

from maivn_tools import WorkbooksToolSet

if TYPE_CHECKING:
    from maivn_tools.connectors.workbooks.models import (
        ExcelNumberFormat,
        RangeFormat,
        WorkbookCompositionContent,
        WorkbookHandle,
    )

_INITIAL_VOLUME = 100
_REVISED_VOLUME = 120


def _expected_formulas() -> dict[str, str]:
    """Return the expected calculation graph independently of rendered cells."""
    expected = {
        'Forecast!B4': '=Assumptions!$B$2',
        'Forecast!B5': '=B4*(1+Assumptions!$B$3)',
        'Forecast!B6': '=B5*(1+Assumptions!$B$3)',
        'Forecast!C8': '=SUM(C4:C6)',
        'Forecast!D8': '=SUM(D4:D6)',
        'Forecast!E8': '=SUM(E4:E6)',
        'Summary!B4': '=Forecast!C8',
        'Summary!B5': '=Forecast!D8',
        'Summary!B6': '=Forecast!E8',
    }
    for row in range(4, 7):
        expected.update(
            {
                f'Forecast!C{row}': f'=B{row}*Assumptions!$B$4',
                f'Forecast!D{row}': f'=C{row}*Assumptions!$B$5+Assumptions!$B$6',
                f'Forecast!E{row}': f'=C{row}-D{row}',
                f'Forecast!F{row}': f'=E{row}/C{row}',
                f'Summary!B{row + 6}': f'=Forecast!C{row}',
                f'Summary!C{row + 6}': f'=Forecast!E{row}',
            }
        )
    return expected


def _place(  # noqa: PLR0913 - keep fixture placement and formatting beside its content
    tools: WorkbooksToolSet,
    handle: WorkbookHandle,
    sheet: str,
    name: str,
    cell: str,
    content: WorkbookCompositionContent,
    formatting: RangeFormat,
) -> None:
    reference = tools.compose_artifact(handle, name, content)
    tools.put_range(handle, name, sheet, cell, reference, formatting)


def run(output_dir: Path) -> dict[str, object]:
    """Build and revise a three-sheet forecast; leave calculation to an Office engine."""
    output_dir.mkdir(parents=True, exist_ok=True)
    tools = WorkbooksToolSet(output_dir / 'workspace')
    handle = tools.create_workbook('harbor-forecast.xlsx')
    for sheet, name in [
        ('summary', 'Summary'),
        ('assumptions', 'Assumptions'),
        ('forecast', 'Forecast'),
    ]:
        tools.put_sheet(
            handle,
            sheet,
            name,
            {'anchor': 'end'},
            freeze_panes='A4',
            print_settings={'mode': 'fit_width', 'orientation': 'landscape'},
        )
    _place(
        tools,
        handle,
        'assumptions',
        'assumptions-title',
        'A1',
        {'kind': 'row', 'values': ['Harbor service forecast', 'Input value']},
        {'style': 'header_row'},
    )
    _place(
        tools,
        handle,
        'assumptions',
        'assumption-labels',
        'A2',
        {
            'kind': 'matrix',
            'values': [
                ['January subscriptions'],
                ['Monthly growth'],
                ['Price per subscription (USD)'],
                ['Variable cost ratio'],
                ['Monthly fixed cost (USD)'],
            ],
        },
        {},
    )
    assumptions: list[tuple[int, str, float, ExcelNumberFormat]] = [
        (2, 'starting-volume', 100, '#,##0'),
        (3, 'growth', 0.1, '0.0%'),
        (4, 'price', 100, '$#,##0.00'),
        (5, 'cost-ratio', 0.4, '0.0%'),
        (6, 'fixed-cost', 2000, '$#,##0.00'),
    ]
    for row, name, value, number_format in assumptions:
        _place(
            tools,
            handle,
            'assumptions',
            name,
            f'B{row}',
            {'kind': 'scalar', 'value': value},
            {'style': 'input', 'number_format': number_format},
        )
    _place(
        tools,
        handle,
        'assumptions',
        'assumptions-note',
        'A9',
        {
            'kind': 'row',
            'values': [
                'Illustrative planning data. Fractional subscriptions represent expected '
                'volume, not actual customer counts. Edit blue input cells to recalculate '
                'the forecast.'
            ],
        },
        {},
    )

    _place(
        tools,
        handle,
        'forecast',
        'forecast-title',
        'A1',
        {'kind': 'row', 'values': ['Monthly forecast · Q1 2027']},
        {'style': 'accent'},
    )
    _place(
        tools,
        handle,
        'forecast',
        'forecast-headers',
        'A3',
        {
            'kind': 'row',
            'values': [
                'Month',
                'Subscriptions',
                'Revenue (USD)',
                'Costs (USD)',
                'Profit (USD)',
                'Margin',
            ],
        },
        {'style': 'header_row'},
    )
    _place(
        tools,
        handle,
        'forecast',
        'months',
        'A4',
        {'kind': 'matrix', 'values': [['Jan'], ['Feb'], ['Mar']]},
        {},
    )
    _place(
        tools,
        handle,
        'forecast',
        'volumes',
        'B4',
        {
            'kind': 'formulas',
            'formulas': [
                ['=Assumptions!$B$2'],
                ['=B4*(1+Assumptions!$B$3)'],
                ['=B5*(1+Assumptions!$B$3)'],
            ],
        },
        {'number_format': '#,##0.00'},
    )
    _place(
        tools,
        handle,
        'forecast',
        'calculated-money',
        'C4',
        {
            'kind': 'formulas',
            'formulas': [
                [
                    f'=B{row}*Assumptions!$B$4',
                    f'=C{row}*Assumptions!$B$5+Assumptions!$B$6',
                    f'=C{row}-D{row}',
                ]
                for row in range(4, 7)
            ],
        },
        {'number_format': '$#,##0.00'},
    )
    _place(
        tools,
        handle,
        'forecast',
        'margins',
        'F4',
        {'kind': 'formulas', 'formulas': [[f'=E{row}/C{row}'] for row in range(4, 7)]},
        {'number_format': '0.0%'},
    )
    _place(
        tools,
        handle,
        'forecast',
        'total-label',
        'A8',
        {'kind': 'row', 'values': ['Quarter total']},
        {'style': 'total'},
    )
    _place(
        tools,
        handle,
        'forecast',
        'quarter-money',
        'C8',
        {'kind': 'formulas', 'formulas': [['=SUM(C4:C6)', '=SUM(D4:D6)', '=SUM(E4:E6)']]},
        {'style': 'total', 'number_format': '$#,##0.00'},
    )

    _place(
        tools,
        handle,
        'summary',
        'summary-title',
        'A1',
        {'kind': 'row', 'values': ['Harbor service forecast · Q1 2027']},
        {'style': 'accent'},
    )
    _place(
        tools,
        handle,
        'summary',
        'summary-headers',
        'A3',
        {'kind': 'row', 'values': ['Quarter metric', 'Amount (USD)']},
        {'style': 'header_row'},
    )
    _place(
        tools,
        handle,
        'summary',
        'summary-labels',
        'A4',
        {'kind': 'matrix', 'values': [['Revenue'], ['Costs'], ['Profit']]},
        {},
    )
    _place(
        tools,
        handle,
        'summary',
        'summary-values',
        'B4',
        {'kind': 'formulas', 'formulas': [['=Forecast!C8'], ['=Forecast!D8'], ['=Forecast!E8']]},
        {'number_format': '$#,##0.00'},
    )
    _place(
        tools,
        handle,
        'summary',
        'chart-headers',
        'A9',
        {'kind': 'row', 'values': ['Month', 'Revenue', 'Profit']},
        {
            'style': 'header_row',
            'chart': {
                'type': 'bar',
                'title': 'Monthly revenue and profit (USD)',
                'anchor': 'E3',
                'source_range': 'A9:C12',
            },
        },
    )
    _place(
        tools,
        handle,
        'summary',
        'chart-months',
        'A10',
        {'kind': 'matrix', 'values': [['Jan'], ['Feb'], ['Mar']]},
        {},
    )
    _place(
        tools,
        handle,
        'summary',
        'chart-values',
        'B10',
        {
            'kind': 'formulas',
            'formulas': [[f'=Forecast!C{row}', f'=Forecast!E{row}'] for row in range(4, 7)],
        },
        {'number_format': '$#,##0.00'},
    )

    initial = tools.render_workbook(handle)
    initial_bytes = Path(initial.path).read_bytes()
    initial_path = output_dir / 'forecast-initial.xlsx'
    initial_path.write_bytes(initial_bytes)
    before = load_workbook(initial.path)
    reopened = WorkbooksToolSet(output_dir / 'workspace')
    reopened.compose_artifact(handle, 'starting-volume', {'kind': 'scalar', 'value': 120})
    revised = reopened.render_workbook(handle, overwrite=True)
    revised_path = output_dir / 'forecast-revised.xlsx'
    revised_path.write_bytes(Path(revised.path).read_bytes())
    after = load_workbook(revised_path)
    checks: dict[str, bool] = {
        'three_sheets': after.sheetnames == ['Summary', 'Assumptions', 'Forecast'],
        'input_edited': before['Assumptions']['B2'].value == _INITIAL_VOLUME
        and after['Assumptions']['B2'].value == _REVISED_VOLUME,
        'formulas_preserved': all(
            {
                f'{sheet.title}!{cell.coordinate}': cell.value
                for sheet in book
                for row in sheet
                for cell in row
                if cell.data_type == 'f'
            }
            == _expected_formulas()
            for book in (before, after)
        ),
        'formats_preserved': all(
            before[sheet.title][cell.coordinate].number_format == cell.number_format
            and before[sheet.title][cell.coordinate].style_id == cell.style_id
            for sheet in after
            for row in sheet
            for cell in row
        ),
        'freeze_preserved': all(sheet.freeze_panes == 'A4' for sheet in after),
        'recalculate_on_open': bool(after.calculation and after.calculation.fullCalcOnLoad),
        'prior_bytes_unchanged': Path(initial.path).read_bytes() == initial_bytes,
        'changed_bytes': initial.sha256 != revised.sha256,
    }
    with ZipFile(revised_path) as package:
        checks['chart_linked_to_formula_cells'] = b"'Summary'!$B$10:$B$12" in package.read(
            'xl/charts/chart1.xml'
        )
    if not all(checks.values()):
        raise AssertionError(checks)
    return {
        'format': 'xlsx',
        'initial': str(initial_path),
        'revised': str(revised_path),
        'checks': checks,
        'sha256': hashlib.sha256(revised_path.read_bytes()).hexdigest(),
        'calculation_expectations': {
            'initial': {'Summary!B4': 33100, 'Summary!B5': 19240, 'Summary!B6': 13860},
            'revised': {'Summary!B4': 39720, 'Summary!B5': 21888, 'Summary!B6': 17832},
        },
        'visual_review': 'pending',
        'limitations': [
            'Calculation expectations are independent arithmetic; an Office engine must '
            'verify cached results.',
            'Workbook checks cover the fixture, not arbitrary financial models.',
        ],
        'formula_count': sum(
            cell.data_type == 'f' for sheet in after for row in sheet for cell in row
        ),
        'note_height': cast('Any', after['Assumptions'].row_dimensions[9]).height,
    }
