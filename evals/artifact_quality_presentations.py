"""Deterministic executive deck fixture through the public presentation tools."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import TYPE_CHECKING, Any, cast

from pptx import Presentation

from maivn_tools import PresentationsToolSet

if TYPE_CHECKING:
    from maivn_tools.connectors.presentations.models import (
        PresentationCompositionContent,
        PresentationElement,
        PresentationHandle,
    )

_SLIDE_COUNT = 5


def _put(  # noqa: PLR0913 - fixture placement keeps geometry next to its content
    tools: PresentationsToolSet,
    deck: PresentationHandle,
    slide: str,
    name: str,
    kind: str,
    content: PresentationCompositionContent,
    y: float,
    height: float,
    font_size: float = 20,
) -> None:
    reference = tools.compose_artifact(deck, name, content)
    tools.put_element(
        deck,
        slide,
        name,
        cast(
            'PresentationElement',
            {
                'type': kind,
                'composition': reference,
                'x': 0.7,
                'y': y,
                'width': 11.9,
                'height': height,
                'font_size': font_size,
                'font_family': 'Arial',
            },
        ),
    )


def _build(tools: PresentationsToolSet) -> PresentationHandle:
    deck = tools.create_presentation('executive-review.pptx')
    titles = {
        'summary': 'September operating review',
        'metrics': 'Resolution time fell 30% during the pilot',
        'delivery': 'Rollout responsibilities and decision gates',
        'risks': 'Coverage limits the next expansion',
        'decision': 'A staged expansion keeps the decision reversible',
    }
    for slide, title in titles.items():
        tools.put_slide(deck, slide, {'anchor': 'end'})
        _put(
            tools,
            deck,
            slide,
            f'{slide}-title',
            'title',
            {'kind': 'text', 'text': title},
            0.5,
            1.35,
            32,
        )
    _put(
        tools,
        deck,
        'summary',
        'summary-body',
        'text',
        {
            'kind': 'text',
            'text': 'The support pilot reduced median resolution time from 10 to 7 hours. '
            'Customer satisfaction held above the 90% target.\n\n'
            'The team recommends extending the pilot to one additional region, with a coverage '
            'review before any broader rollout.',
        },
        2.1,
        3.2,
        24,
    )
    _put(
        tools,
        deck,
        'summary',
        'source',
        'text',
        {
            'kind': 'text',
            'text': 'Illustrative benchmark data. Prepared for the operating committee.',
        },
        6.2,
        0.55,
        17,
    )
    _put(
        tools,
        deck,
        'metrics',
        'metrics-chart',
        'chart',
        {
            'kind': 'chart',
            'categories': ['June', 'July', 'August'],
            'series': [
                {'name': 'Median resolution (hours)', 'values': [10.0, 8.5, 7.0]},
                {'name': 'Target (hours)', 'values': [8.0, 8.0, 8.0]},
            ],
        },
        2.0,
        3.8,
        18,
    )
    _put(
        tools,
        deck,
        'metrics',
        'metrics-note',
        'text',
        {'kind': 'text', 'text': 'August: 7.0 hours versus an 8.0-hour target. Satisfaction: 92%.'},
        6.05,
        0.7,
        20,
    )
    _put(
        tools,
        deck,
        'delivery',
        'delivery-table',
        'table',
        {
            'kind': 'table',
            'rows': [
                ['Workstream', 'Owner', 'Due', 'Decision gate'],
                ['Regional readiness', 'Operations', '18 September', 'Weekend coverage confirmed'],
                [
                    'Support training',
                    'Enablement',
                    '25 September',
                    'All shift leads complete the simulation',
                ],
                [
                    'Pilot extension',
                    'Regional lead',
                    '2 October',
                    'Seven days within the response target',
                ],
            ],
        },
        2.0,
        3.5,
        18,
    )
    _put(
        tools,
        deck,
        'delivery',
        'delivery-note',
        'text',
        {
            'kind': 'text',
            'text': 'The regional lead owns the launch decision after both readiness gates pass.',
        },
        5.9,
        0.8,
        20,
    )
    _put(
        tools,
        deck,
        'risks',
        'risk-body',
        'list',
        {
            'kind': 'list',
            'items': [
                'Weekend staffing remains below the weekday baseline. '
                'Operations will close the gap before launch.',
                'The pilot covers one region. A second region will test '
                'whether the process generalizes.',
                'Pause expansion if median resolution exceeds eight hours '
                'for two consecutive weeks.',
            ],
        },
        2.1,
        4.3,
        24,
    )
    _put(
        tools,
        deck,
        'decision',
        'recommendation',
        'text',
        {
            'kind': 'text',
            'text': 'Approve a four-week extension in one additional region.\n\n'
            'Cap incremental staffing at two specialists. Review resolution time and satisfaction '
            'each Friday. The operating committee will decide on broader rollout after the trial.',
        },
        2.1,
        3.5,
        24,
    )
    return deck


def _facts(path: Path) -> dict[str, object]:
    deck = Presentation(str(path))
    text: list[str] = []
    charts: list[list[list[float]]] = []
    tables: list[list[list[str]]] = []
    for slide in deck.slides:
        for shape in slide.shapes:
            item = cast('Any', shape)
            if item.has_text_frame:
                text.append(str(item.text))
            if item.has_chart:
                charts.append([list(series.values) for series in item.chart.series])
            if item.has_table:
                tables.append([[cell.text for cell in row.cells] for row in item.table.rows])
    return {'slide_count': len(deck.slides), 'text': text, 'charts': charts, 'tables': tables}


def run(output_dir: Path) -> dict[str, object]:
    """Build original and revised files; report objective content and layout checks."""
    output_dir = output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    tools = PresentationsToolSet(output_dir / 'workspace')
    deck = _build(tools)
    initial = tools.render_presentation(deck, overwrite=True)
    initial_path = output_dir / 'initial.pptx'
    initial_path.write_bytes(Path(initial.path).read_bytes())
    initial_issues = tools.inspect_layout(deck)
    before = tools.read_presentation(deck)
    tools.compose_artifact(
        deck,
        'recommendation',
        {
            'kind': 'text',
            'text': 'Approve a six-week extension in one additional region.\n\n'
            'Cap incremental staffing at two specialists. Review resolution time and satisfaction '
            'each Friday. The operating committee will decide on broader rollout after the trial.',
        },
    )
    revised = tools.render_presentation(deck, overwrite=True)
    revised_path = output_dir / 'revised.pptx'
    revised_path.write_bytes(Path(revised.path).read_bytes())
    after = tools.read_presentation(deck)
    initial_facts, revised_facts = _facts(initial_path), _facts(revised_path)
    initial_text = '\n'.join(cast('list[str]', initial_facts['text']))
    revised_text = '\n'.join(cast('list[str]', revised_facts['text']))
    return {
        'initial_path': str(initial_path),
        'revised_path': str(revised_path),
        'initial_sha256': hashlib.sha256(initial_path.read_bytes()).hexdigest(),
        'revised_sha256': hashlib.sha256(revised_path.read_bytes()).hexdigest(),
        'initial_layout_issues': initial_issues,
        'revised_layout_issues': tools.inspect_layout(deck),
        'initial_facts': initial_facts,
        'revised_facts': revised_facts,
        'checks': {
            'five_slides': initial_facts['slide_count']
            == revised_facts['slide_count']
            == _SLIDE_COUNT,
            'chart_values': initial_facts['charts'] == [[[10.0, 8.5, 7.0], [8.0, 8.0, 8.0]]],
            'chart_preserved': initial_facts['charts'] == revised_facts['charts'],
            'complete_table': initial_facts['tables']
            == [
                [
                    ['Workstream', 'Owner', 'Due', 'Decision gate'],
                    [
                        'Regional readiness',
                        'Operations',
                        '18 September',
                        'Weekend coverage confirmed',
                    ],
                    [
                        'Support training',
                        'Enablement',
                        '25 September',
                        'All shift leads complete the simulation',
                    ],
                    [
                        'Pilot extension',
                        'Regional lead',
                        '2 October',
                        'Seven days within the response target',
                    ],
                ]
            ],
            'table_preserved': initial_facts['tables'] == revised_facts['tables'],
            'recommendation_revised': 'Approve a four-week extension' in initial_text
            and 'Approve a six-week extension' in revised_text
            and 'Approve a four-week extension' not in revised_text,
            'risk_content_present': 'two consecutive weeks' in initial_text
            and 'two consecutive weeks' in revised_text,
            'no_estimated_overflow': not initial_issues and not tools.inspect_layout(deck),
            'slide_geometry_preserved': before['slides'] == after['slides'],
            'revision_changed': initial.sha256 != revised.sha256,
            'only_recommendation_changed': [
                c for c in before['compositions'] if c['composition_id'] != 'recommendation'
            ]
            == [c for c in after['compositions'] if c['composition_id'] != 'recommendation'],
            'initial_file_immutable': Path(initial.path).read_bytes() == initial_path.read_bytes(),
        },
        'limitations': [
            'Synthetic data; no claim about business performance.',
            'Layout checks are estimates. Office rendering and visual review are separate gates.',
        ],
    }
